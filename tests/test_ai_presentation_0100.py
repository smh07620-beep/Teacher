from __future__ import annotations

import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from flask import Flask

from teacher_app.materials import ai_presentation_jobs
from teacher_app.materials import ai_presentation_repository as presentation_repository
from teacher_app.materials import ai_presentation_runtime
from teacher_app.materials import ai_presentation_routes
from teacher_app.materials.ai_presentation_routes import _artifact_ready, _presentation_allowed
from teacher_app.materials.ai_presentation_repository import publication_receipt_key
from teacher_app.materials.ai_presentation_repository import sanitize_provenance
from teacher_app.materials.ai_presentation_runtime import (
    normalize_slides,
    parse_slide_outline,
    render_pptx,
)
from teacher_app.materials.ai_presentation_storage import PPTX_MIME, PresentationStorage


def _presentation_migration():
    """Import migrations only while the test runs, not during unittest discovery.

    The global migration registry is order-sensitive.  Importing 0099/0100 at
    module discovery time would register them before older additive migration
    modules are loaded through the canonical production compatibility chain.
    """
    from teacher_app.maintenance import ai_presentation_migration

    return ai_presentation_migration


class AiPresentationMigrationTests(unittest.TestCase):
    def test_fresh_0099_then_0100_creates_required_tables(self):
        migration = _presentation_migration()
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        migration.ai_presentations_99(conn, "sqlite")
        migration.ai_presentation_production_hardening_100(conn, "sqlite")
        tables = {
            row[0]
            for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
        }
        self.assertIn("ai_presentation_templates", tables)
        self.assertIn("ai_presentation_jobs", tables)
        self.assertIn("ai_presentations", tables)
        self.assertIn("ai_presentation_publications", tables)

    def test_0100_upgrades_local_only_0099_additively(self):
        migration = _presentation_migration()
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        conn.execute(
            "CREATE TABLE ai_presentation_templates ("
            "id TEXT PRIMARY KEY,name TEXT NOT NULL,group_key TEXT NOT NULL,training_area TEXT NOT NULL)"
        )
        conn.execute(
            "CREATE TABLE ai_presentations ("
            "id TEXT PRIMARY KEY,title TEXT NOT NULL,group_key TEXT NOT NULL,training_area TEXT NOT NULL)"
        )
        conn.execute(
            "INSERT INTO ai_presentations(id,title,group_key,training_area) VALUES(?,?,?,?)",
            ("legacy-ppt", "Legacy", "grpBio", "internal"),
        )
        migration.ai_presentation_production_hardening_100(conn, "sqlite")
        columns = {
            row[1] for row in conn.execute("PRAGMA table_info(ai_presentations)").fetchall()
        }
        self.assertTrue(
            {
                "presentation_family_id",
                "parent_version_id",
                "revision_number",
                "artifact_backend",
                "artifact_storage_key",
                "artifact_sha256",
                "artifact_bytes",
                "artifact_mime_type",
                "published_at",
            }.issubset(columns)
        )
        row = conn.execute("SELECT id,title FROM ai_presentations WHERE id='legacy-ppt'").fetchone()
        self.assertEqual((row["id"], row["title"]), ("legacy-ppt", "Legacy"))

    def test_0102_adds_persisted_provenance_without_replacing_rows(self):
        migration = _presentation_migration()
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        migration.ai_presentations_99(conn, "sqlite")
        migration.ai_presentation_production_hardening_100(conn, "sqlite")
        conn.execute("INSERT INTO ai_presentations(id,material_id,draft_id,group_key,training_area,title,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", ("ppt-old", "mat-1", "draft-1", "grpBio", "internal", "Old", "now", "now"))
        # Preserve global migration registration order even when this focused
        # test runs before the AI-video test module.
        from teacher_app.maintenance import ai_video_migration as _ai_video_migration  # noqa: F401
        from teacher_app.maintenance.ai_presentation_phase2_migration import ai_presentation_provenance_102
        ai_presentation_provenance_102(conn, "sqlite")
        columns = {row[1] for row in conn.execute("PRAGMA table_info(ai_presentations)").fetchall()}
        self.assertIn("provenance_json", columns)
        self.assertEqual(conn.execute("SELECT title FROM ai_presentations WHERE id='ppt-old'").fetchone()[0], "Old")

    def test_0103_adds_layout_profile_and_publication_snapshot_additively(self):
        migration = _presentation_migration()
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        migration.ai_presentations_99(conn, "sqlite")
        migration.ai_presentation_production_hardening_100(conn, "sqlite")
        from teacher_app.maintenance.ai_presentation_phase3_migration import ai_presentation_publishing_workflow_103
        ai_presentation_publishing_workflow_103(conn, "sqlite")
        template_columns = {row[1] for row in conn.execute("PRAGMA table_info(ai_presentation_templates)").fetchall()}
        publication_columns = {row[1] for row in conn.execute("PRAGMA table_info(ai_presentation_publications)").fetchall()}
        self.assertIn("layout_profile_json", template_columns)
        self.assertTrue({"presentation_family_id", "presentation_revision_number", "snapshot_json"}.issubset(publication_columns))


class AiPresentationRuntimeTests(unittest.TestCase):
    def test_numbered_bullets_are_not_misread_as_slide_headers(self):
        slides = parse_slide_outline(
            "第 1 張：CBC 概念\n1. WBC\n2. RBC\n第 2 張：判讀\n- 先確認 QC\n- 再判讀結果"
        )
        self.assertEqual(len(slides), 2)
        self.assertEqual(slides[0]["title"], "CBC 概念")
        self.assertEqual(slides[0]["bullets"], ["WBC", "RBC"])

    def test_teacher_slide_controls_normalize_order_and_enabled(self):
        slides = normalize_slides(
            [
                {"id": "b", "order": 9, "enabled": False, "title": "B", "bullets": ["x"]},
                {"id": "a", "order": 2, "enabled": True, "title": "A", "bullets": ["y"]},
            ]
        )
        self.assertEqual([item["order"] for item in slides], [1, 2])
        self.assertFalse(slides[0]["enabled"])
        self.assertTrue(slides[1]["enabled"])

    def test_layout_and_safe_content_blocks_are_normalized_without_urls(self):
        slides = normalize_slides([{
            "layout": "comparison", "title": "方法比較", "bullets": [],
            "blocks": [
                {"type": "comparison", "leftTitle": "A", "rightTitle": "B", "leftItems": ["快"], "rightItems": ["準"]},
                {"type": "chart", "labels": ["一", "二"], "values": [1, 2]},
                {"type": "image", "url": "https://not-allowed.example/a.png", "altText": "缺圖"},
            ],
        }])
        self.assertEqual(slides[0]["layout"], "comparison")
        self.assertEqual([block["type"] for block in slides[0]["blocks"]], ["comparison", "chart", "image"])
        self.assertNotIn("url", slides[0]["blocks"][2])

    def test_layout_profile_is_allowlisted(self):
        profile = presentation_repository.sanitize_layout_profile({"layoutMap": {"content": "Title and Content", "evil": "ignored"}})
        self.assertEqual(profile, {"layoutMap": {"content": "Title and Content"}})

    @unittest.skipIf(ai_presentation_runtime.Presentation is None, "python-pptx is installed only on the AI Worker")
    def test_rendered_pptx_contains_allowlisted_provenance(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "test.pptx"
            render_pptx(
                title="AI Test",
                slides=[{"title": "Slide", "bullets": ["Point"], "enabled": True}],
                output_path=output,
                provenance={
                    "sourceMaterialId": "mat-1",
                    "sourceDraftId": "draft-1",
                    "sourceJobId": "job-1",
                    "sourceChunkIds": ["chunk-1"],
                    "teacherApprovedBy": "teacher-a",
                    "teacherApprovedAt": "2026-09-30T00:00:00+00:00",
                },
            )
            self.assertGreater(output.stat().st_size, 1024)
            from pptx import Presentation

            rendered = Presentation(str(output))
            self.assertLessEqual(len(rendered.core_properties.comments), 255)
            self.assertIn("mat-1", rendered.core_properties.comments)
            self.assertNotIn("token=", rendered.core_properties.comments.lower())
            notes = "\n".join(
                slide.notes_slide.notes_text_frame.text
                for slide in rendered.slides
                if slide.has_notes_slide
            )
            self.assertIn('"sourceChunkIds":["chunk-1"]', notes)
            self.assertIn('"teacherApprovedBy":"teacher-a"', notes)

    @unittest.skipIf(ai_presentation_runtime.Presentation is None, "python-pptx is installed only on the AI Worker")
    def test_renderer_handles_quality_blocks_and_missing_image_as_fallback(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "blocks.pptx"
            render_pptx(
                title="Blocks", output_path=output, provenance={"sourceMaterialId": "mat-1"},
                slides=normalize_slides([{"title": "比較", "layout": "comparison", "blocks": [
                    {"type": "comparison", "leftTitle": "舊流程", "rightTitle": "新流程", "leftItems": ["慢"], "rightItems": ["快"]},
                    {"type": "table", "headers": ["項目", "值"], "rows": [["QC", "Pass"]]},
                    {"type": "chart", "labels": ["一", "二"], "values": [1, 2]},
                    {"type": "image", "altText": "尚未上傳的圖"},
                ]}]),
            )
            self.assertGreater(output.stat().st_size, 1024)

    def test_provenance_rejects_secret_like_allowlisted_value(self):
        with self.assertRaises(ValueError):
            ai_presentation_runtime._provenance(
                {"sourceMaterialId": "mat-1", "sourceJobId": "token=do-not-embed"}
            )

    def test_provenance_rejects_local_filesystem_paths(self):
        for path in (r"C:\Users\worker\presentation.pptx", "/home/worker/presentation.pptx"):
            with self.subTest(path=path), self.assertRaises(ValueError):
                ai_presentation_runtime._provenance({"sourceMaterialId": path})

    def test_persisted_provenance_is_allowlisted_and_rejects_secrets(self):
        result = sanitize_provenance({"sourceMaterialId": "mat-1", "sourceChunkIds": ["chunk-1"], "ignored": "nope"})
        self.assertEqual(result["sourceMaterialId"], "mat-1")
        self.assertEqual(result["sourceChunkIds"], ["chunk-1"])
        self.assertNotIn("ignored", result)
        with self.assertRaisesRegex(ValueError, "敏感資訊"):
            sanitize_provenance({"sourceJobId": "token=do-not-store"})

    def test_provenance_includes_bounded_revision_identity(self):
        self.assertEqual(sanitize_provenance({"revisionNumber": 7})["revisionNumber"], 7)

    def test_revision_worker_rechecks_group_and_area_scope(self):
        current = {
            "id": "ppt-r2",
            "status": "draft",
            "artifactStorageKey": "",
            "artifactSha256": "",
            "artifactBytes": 0,
            "group": "grpBio",
            "area": "internal",
            "draftId": "draft-1",
            "materialId": "mat-1",
        }
        job = {
            "group": "grpBio",
            "area": "internal",
            "request": {"presentationId": "ppt-r2", "renderRevision": True},
        }
        draft = {"id": "draft-1", "group": "grpBio", "area": "pgy"}
        source = {"id": "mat-1", "group": "grpBio", "area": "internal"}
        with (
            patch.object(ai_presentation_runtime.repository, "get_presentation", return_value=current),
            patch.object(ai_presentation_runtime.media_script_repository, "get_script", return_value=draft),
            patch.object(ai_presentation_runtime.material_repository, "get_material", return_value=source),
        ):
            with self.assertRaisesRegex(RuntimeError, "授權範圍已改變"):
                ai_presentation_runtime.generate_revision(job=job, storage=object())

    def test_local_storage_is_fail_closed_without_explicit_development_opt_in(self):
        with patch.dict(
            os.environ,
            {
                "AI_PRESENTATION_STORAGE_BACKEND": "local",
                "AI_PRESENTATION_ALLOW_LOCAL_STORAGE": "false",
            },
            clear=False,
        ):
            with self.assertRaises(RuntimeError):
                PresentationStorage().backend()


class AiPresentationRevisionTests(unittest.TestCase):
    def test_structural_edit_creates_new_revision_without_stale_artifact(self):
        current = {
            "id": "ppt-r1",
            "presentationFamilyId": "ppt-family",
            "revisionNumber": 1,
            "materialId": "mat-1",
            "draftId": "draft-1",
            "templateId": "tpl-1",
            "group": "grpBio",
            "area": "internal",
            "title": "Original",
            "slides": [{"title": "A", "bullets": ["one"]}],
            "sourceJobId": "pptjob-1",
            "provider": "groq",
            "model": "model-1",
            "artifactBackend": "r2",
            "artifactStorageKey": "old.pptx",
            "artifactSha256": "a" * 64,
            "artifactBytes": 1234,
            "artifactMimeType": PPTX_MIME,
        }
        expected = {"id": "ppt-r2"}
        with (
            patch.object(presentation_repository, "next_revision_number", return_value=2),
            patch.object(presentation_repository, "create_presentation", return_value=expected) as create,
        ):
            result = presentation_repository.create_revision(
                current,
                actor_username="teacher-a",
                title="Edited",
                slides=[{"title": "B", "bullets": ["two"]}],
            )
        self.assertEqual(result, expected)
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["presentation_family_id"], "ppt-family")
        self.assertEqual(kwargs["parent_version_id"], "ppt-r1")
        self.assertEqual(kwargs["revision_number"], 2)
        self.assertEqual(kwargs["artifact"], {})

    def test_artifact_metadata_must_be_complete_before_attaching(self):
        with self.assertRaises(ValueError):
            presentation_repository.update_presentation_artifact(
                "ppt-r2",
                artifact={
                    "backend": "r2",
                    "key": "ai-presentations/artifacts/ppt-r2/file.pptx",
                    "sha256": "bad",
                    "byteSize": 123,
                    "mimeType": PPTX_MIME,
                },
                actor_username="teacher-a",
            )

    def test_artifact_ready_requires_shared_provider_and_complete_metadata(self):
        artifact = {
            "artifactBackend": "r2",
            "artifactStorageKey": "ai-presentations/artifacts/ppt-r2/file.pptx",
            "artifactSha256": "a" * 64,
            "artifactBytes": 4096,
            "artifactMimeType": PPTX_MIME,
        }
        self.assertTrue(_artifact_ready(artifact))
        with patch.dict(os.environ, {"AI_PRESENTATION_ALLOW_LOCAL_STORAGE": "false"}, clear=False):
            self.assertFalse(_artifact_ready({**artifact, "artifactBackend": "local"}))


class AiPresentationPublicationTests(unittest.TestCase):
    def test_publication_receipt_key_is_deterministic(self):
        first = publication_receipt_key("ppt-1", "mat-1")
        second = publication_receipt_key("ppt-1", "mat-1")
        self.assertEqual(first, second)
        self.assertNotEqual(first, publication_receipt_key("ppt-1", "mat-2"))

    def test_replayed_publication_returns_existing_receipt_without_new_write(self):
        existing = {
            "id": "pub-1",
            "presentationId": "ppt-1",
            "publicationMaterialId": "mat-1",
            "receiptKey": publication_receipt_key("ppt-1", "mat-1"),
        }
        with (
            patch.object(presentation_repository, "get_publication_by_key", return_value=existing),
            patch.object(presentation_repository.common_db, "transaction") as transaction,
        ):
            result = presentation_repository.create_publication(
                presentation_id="ppt-1",
                publication_material_id="mat-1",
                actor_username="teacher-a",
                receipt={"presentationSha256": "a" * 64},
            )
        self.assertEqual(result, existing)
        transaction.assert_not_called()


class AiPresentationWorkerQueueTests(unittest.TestCase):
    def test_revision_job_dispatches_to_worker_revision_renderer(self):
        queued = {"id": "pptjob-1"}
        claimed = {
            "id": "pptjob-1",
            "request": {"presentationId": "ppt-r2", "renderRevision": True},
        }
        processor = ai_presentation_jobs.AiPresentationJobProcessor()
        with (
            patch.object(presentation_repository, "list_queued", return_value=[queued]),
            patch.object(presentation_repository, "claim_job", return_value=claimed),
            patch.object(presentation_repository, "complete_job", return_value=True),
            patch.object(ai_presentation_jobs.ai_presentation_runtime, "generate_revision", return_value={"presentationId": "ppt-r2"}) as revision,
            patch.object(ai_presentation_jobs.ai_presentation_runtime, "generate_presentation") as initial,
        ):
            self.assertTrue(processor.run_next_queued())
        revision.assert_called_once()
        initial.assert_not_called()

    def test_failed_job_retry_reuses_the_same_idempotency_job(self):
        failed = {"id": "pptjob-failed", "status": "failed", "attempts": 1}
        with patch.object(presentation_repository, "retry_failed_job", return_value={**failed, "status": "queued"}) as retry:
            result = ai_presentation_jobs.retry(job=failed)
        self.assertEqual(result["id"], "pptjob-failed")
        retry.assert_called_once()


class AiPresentationRbacTests(unittest.TestCase):
    def test_system_admin_can_operate_but_cannot_teacher_approve(self):
        user = {"role": "system_admin"}
        self.assertTrue(_presentation_allowed(user, "presentation.create"))
        self.assertTrue(_presentation_allowed(user, "presentation.edit"))
        self.assertFalse(_presentation_allowed(user, "presentation.approve"))
        self.assertTrue(_presentation_allowed(user, "presentation.publish"))


class AiPresentationProvenanceRouteTests(unittest.TestCase):
    def test_scoped_teacher_can_read_allowlisted_persisted_provenance(self):
        app = Flask(__name__)
        owner = SimpleNamespace(
            app=app,
            _current_user=lambda: {"username": "teacher-a", "role": "clinical_teacher"},
        )
        ai_presentation_routes.register_ai_presentation_routes(owner)
        revision = {
            "id": "ppt-r2", "presentationFamilyId": "ppt-family", "revisionNumber": 2,
            "group": "grpBio", "provenance": {
                "sourceMaterialId": "mat-1", "sourceDraftId": "draft-1", "sourceJobId": "job-1",
                "sourceChunkIds": ["chunk-1"], "provider": "local", "model": "test-model",
                "templateId": "", "teacherApprovedBy": "teacher-a", "teacherApprovedAt": "2026-09-30T00:00:00+00:00",
            },
        }
        with patch.object(ai_presentation_routes.repository, "get_presentation", return_value=revision), \
             patch.object(ai_presentation_routes, "_scope", return_value=None):
            response = app.test_client().get("/api/ai-presentations/ppt-r2/provenance")
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["presentationId"], "ppt-r2")
        self.assertEqual(body["provenance"]["sourceChunkIds"], ["chunk-1"])

    def test_teacher_and_group_leader_have_full_presentation_permissions(self):
        for role in ("clinical_teacher", "group_leader"):
            user = {"role": role}
            self.assertTrue(_presentation_allowed(user, "presentation.create"))
            self.assertTrue(_presentation_allowed(user, "presentation.edit"))
            self.assertTrue(_presentation_allowed(user, "presentation.approve"))
            self.assertTrue(_presentation_allowed(user, "presentation.publish"))

    def test_education_admin_cannot_teacher_approve(self):
        user = {"role": "education_admin"}
        self.assertTrue(_presentation_allowed(user, "presentation.create"))
        self.assertTrue(_presentation_allowed(user, "presentation.edit"))
        self.assertFalse(_presentation_allowed(user, "presentation.approve"))
        self.assertTrue(_presentation_allowed(user, "presentation.publish"))


class AiPresentationDeploymentContractTests(unittest.TestCase):
    def test_python_pptx_is_worker_only_dependency(self):
        web = Path("requirements.txt").read_text(encoding="utf-8").lower()
        worker = Path("requirements-ai-worker.txt").read_text(encoding="utf-8").lower()
        self.assertNotIn("python-pptx", web)
        self.assertIn("python-pptx", worker)

    def test_ai_worker_consumes_presentation_queue(self):
        worker = Path("ai_question_worker.py").read_text(encoding="utf-8")
        self.assertIn("AiPresentationJobProcessor", worker)
        self.assertIn("presentation_processor.run_next_queued()", worker)

    def test_phase2_workspace_and_provenance_endpoint_are_registered(self):
        assets = Path("teacher_app/frontend/assets.py").read_text(encoding="utf-8")
        frontend = Path("static/teacher-ai-presentation-1016.js").read_text(encoding="utf-8")
        routes = Path("teacher_app/materials/ai_presentation_routes.py").read_text(encoding="utf-8")
        self.assertIn("0102-ai-presentation-provenance", __import__("release_contract").REQUIRED_MIGRATIONS)
        self.assertIn('/teacher-ai-presentation-1016.js', assets)
        for marker in ('/download', '/provenance', '/reupload', 'immutable revision'):
            self.assertIn(marker, frontend)
        self.assertIn('/api/ai-presentations/<presentation_id>/provenance', routes)

    def test_phase3_contract_keeps_worker_rendering_and_exposes_history(self):
        runtime = Path("teacher_app/materials/ai_presentation_runtime.py").read_text(encoding="utf-8")
        routes = Path("teacher_app/materials/ai_presentation_routes.py").read_text(encoding="utf-8")
        frontend = Path("static/teacher-ai-presentation-1016.js").read_text(encoding="utf-8")
        self.assertIn("0103-ai-presentation-publishing-workflow", __import__("release_contract").REQUIRED_MIGRATIONS)
        self.assertIn("_layout_for", runtime)
        self.assertIn("_render_blocks", runtime)
        self.assertIn("/history", routes)
        self.assertIn("/published/download", routes)
        self.assertIn("版型／區塊編輯", frontend)


if __name__ == "__main__":
    unittest.main()
