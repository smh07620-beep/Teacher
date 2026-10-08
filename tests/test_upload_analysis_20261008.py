"""Upload-time analysis: type decision + Word image summary in one pass."""
import io
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from docx import Document
from docx.shared import Inches

from teacher_app.materials import upload_analysis

ROOT = Path(__file__).resolve().parents[1]


def _png(color):
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), color).save(buf, "PNG")
    buf.seek(0)
    return buf


class UploadAnalysisTests(unittest.TestCase):
    def _docx(self, tmp, images=2):
        doc = Document()
        doc.add_paragraph("血液抹片判讀")
        for index in range(images):
            doc.add_paragraph(f"圖 {index + 1}：中性球形態")
            doc.add_picture(_png((index * 90, 20, 30)), width=Inches(1))
        path = Path(tmp) / "atlas.docx"
        doc.save(path)
        return path

    def test_docx_images_are_counted_with_captions_at_upload_time(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._docx(tmp, images=2)
            result = upload_analysis.build_analysis(
                path, extension=".docx", resolved_type="atlas", method="內容規則判斷", reason="關鍵字分數 12")
        self.assertEqual(result["imageCount"], 2)
        self.assertEqual(result["type"], "atlas")
        self.assertEqual(result["confidence"], "high")
        self.assertFalse(result["needsReview"])
        # the "圖 N：..." label sits in the paragraph before the picture -> section
        self.assertTrue(all("圖" in (item["caption"] + item["section"]) for item in result["images"]))
        # only metadata is persisted, never image bytes
        self.assertNotIn("data", result["images"][0])

    def test_default_classification_is_flagged_for_one_question_review(self):
        result = upload_analysis.build_analysis(
            Path("x.pdf"), extension=".pdf", resolved_type="standard", method="預設分類")
        self.assertEqual(result["confidence"], "low")
        self.assertTrue(result["needsReview"])
        self.assertEqual(result["imageCount"], 0)

    def test_corrupt_docx_never_fails_the_upload(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "bad.docx"
            bad.write_bytes(b"not a zip")
            result = upload_analysis.build_analysis(
                bad, extension=".docx", resolved_type="standard", method="預設分類")
        self.assertEqual(result["imageCount"], 0)
        self.assertIn("imageScanError", result)

    def test_worker_stores_analysis_in_storage_meta(self):
        worker = ROOT.joinpath("material_worker.py").read_text(encoding="utf-8")
        self.assertIn('media_meta["uploadAnalysis"]', worker)
        self.assertIn("upload_analysis.build_analysis", worker)


if __name__ == "__main__":
    unittest.main()
