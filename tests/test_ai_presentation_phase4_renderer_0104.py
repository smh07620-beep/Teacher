from __future__ import annotations

import re
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from teacher_app.materials import ai_presentation_quality as quality
from teacher_app.materials import ai_presentation_repository as repository
from teacher_app.materials import ai_presentation_runtime as runtime


def _slide_text(slide) -> str:
    values: list[str] = []
    for shape in slide.shapes:
        text = str(getattr(shape, "text", "") or "").strip()
        if text:
            values.append(text)
        if bool(getattr(shape, "has_table", False)):
            table = shape.table
            for row in table.rows:
                values.extend(str(cell.text or "") for cell in row.cells)
    return "\n".join(values)


class AiPresentationPhase4RendererTests(unittest.TestCase):
    def test_layout_profile_preserves_legacy_shape_and_allows_safe_placeholders(self):
        legacy = repository.sanitize_layout_profile(
            {"layoutMap": {"content": "Title and Content", "evil": "ignored"}}
        )
        self.assertEqual(legacy, {"layoutMap": {"content": "Title and Content"}})

        profile = repository.sanitize_layout_profile(
            {
                "layoutMap": {"content": "Title and Content"},
                "placeholderMap": {"body": "body", "image": "picture", "evil": "ignored"},
            }
        )
        self.assertEqual(
            profile,
            {
                "layoutMap": {"content": "Title and Content"},
                "placeholderMap": {"body": "body", "image": "picture"},
            },
        )

    def test_runtime_layout_fallback_is_deterministic_and_observable(self):
        layouts = [SimpleNamespace(name="Custom Cover"), SimpleNamespace(name="Custom Body")]
        prs = SimpleNamespace(slide_layouts=layouts)
        manifest = quality.sanitize_quality_manifest({})

        selected = runtime._layout_for(
            prs,
            "table",
            {"layoutMap": {"table": "Missing Layout"}},
            quality_report=manifest,
            slide_id="s-table",
        )

        self.assertIs(selected, layouts[1])
        self.assertEqual(manifest["status"], "warning")
        self.assertIn("TEMPLATE_FALLBACK", {item["code"] for item in manifest["warnings"]})
        self.assertEqual(manifest["warnings"][0]["slideId"], "s-table")

    def test_branding_context_uses_render_date_not_fake_publish_date(self):
        branding = runtime._branding_context(
            title="Phase 4",
            group="生化組",
            area="內部教育訓練",
            teacher="Teacher A",
            revision=4,
        )
        self.assertRegex(branding["renderedDate"], r"^\d{4}-\d{2}-\d{2}$")
        self.assertNotIn("publishedDate", branding)
        self.assertEqual(branding["revisionNumber"], 4)

    @unittest.skipIf(runtime.Presentation is None, "python-pptx is installed only on the AI Worker")
    def test_rendered_pptx_applies_cover_pagination_branding_and_metadata(self):
        slides = [
            {
                "id": "s-text",
                "title": "長內容",
                "bullets": [f"重點 {index}" for index in range(1, 9)],
                "enabled": True,
            },
            {
                "id": "s-table",
                "title": "長表格",
                "enabled": True,
                "blocks": [
                    {
                        "type": "table",
                        "headers": ["項目", "結果"],
                        "rows": [[f"QC-{index}", "Pass"] for index in range(1, 9)],
                    }
                ],
            },
            {
                "id": "s-compare",
                "title": "長比較",
                "enabled": True,
                "blocks": [
                    {
                        "type": "comparison",
                        "leftTitle": "舊流程",
                        "rightTitle": "新流程",
                        "leftItems": [f"舊 {index}" for index in range(1, 12)],
                        "rightItems": [f"新 {index}" for index in range(1, 12)],
                    }
                ],
            },
        ]
        report: dict = {}
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "phase4-render.pptx"
            runtime.render_pptx(
                title="Phase 4 Regression",
                slides=slides,
                output_path=output,
                provenance={"sourceMaterialId": "mat-1", "sourceDraftId": "draft-1"},
                branding={
                    "groupLabel": "生化組",
                    "areaLabel": "內部教育訓練",
                    "teacherName": "Teacher A",
                    "revisionNumber": 3,
                    "renderedDate": "2026-09-30",
                },
                quality_report=report,
            )

            rendered = runtime.Presentation(str(output))
            self.assertEqual(len(rendered.slides), 8)
            self.assertIn("Phase 4 Regression", _slide_text(rendered.slides[0]))
            warning_codes = {item["code"] for item in report["warnings"]}
            self.assertTrue(
                {"TEXT_SPLIT", "TABLE_SPLIT", "COMPARISON_SPLIT"}.issubset(warning_codes)
            )
            rendered_text = "\n".join(_slide_text(slide) for slide in rendered.slides)
            self.assertIn("長內容（續 2）", rendered_text)
            self.assertIn("長表格（續 2）", rendered_text)
            self.assertIn("長比較（續 3）", rendered_text)
            self.assertIn("生化組 · 內部教育訓練 · Teacher A · r3 · 2026-09-30", rendered_text)
            self.assertEqual(
                rendered.core_properties.subject,
                "Teacher AI reviewed teaching presentation",
            )
            self.assertIn(quality.RULESET_VERSION, rendered.core_properties.keywords)
            self.assertEqual(report["slideCount"], 8)
            self.assertEqual(report["errorCount"], 0)

    @unittest.skipIf(runtime.Presentation is None, "python-pptx is installed only on the AI Worker")
    def test_existing_title_layout_is_not_duplicated(self):
        report: dict = {}
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "existing-cover.pptx"
            runtime.render_pptx(
                title="既有封面",
                slides=[
                    {
                        "id": "cover",
                        "title": "既有封面",
                        "layout": "title",
                        "enabled": True,
                    },
                    {
                        "id": "content",
                        "title": "內容",
                        "bullets": ["重點"],
                        "enabled": True,
                    },
                ],
                output_path=output,
                provenance={"sourceMaterialId": "mat-1", "sourceDraftId": "draft-1"},
                quality_report=report,
            )
            rendered = runtime.Presentation(str(output))
            self.assertEqual(len(rendered.slides), 2)
            self.assertEqual(report["slideCount"], 2)

    @unittest.skipIf(
        runtime.Presentation is None or runtime.Image is None,
        "python-pptx and Pillow are installed only on the AI Worker",
    )
    def test_rendered_image_uses_safe_crop_and_visible_caption_source(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.png"
            output = root / "image-render.pptx"
            runtime.Image.new("RGB", (1600, 300)).save(source)
            report: dict = {}

            runtime.render_pptx(
                title="Image Regression",
                slides=[
                    {
                        "id": "s-image",
                        "title": "影像",
                        "layout": "image",
                        "enabled": True,
                        "blocks": [
                            {
                                "type": "image",
                                "asset": {
                                    "backend": "r2",
                                    "key": "ai-presentations/images/test.png",
                                    "sha256": "a" * 64,
                                    "mimeType": "image/png",
                                },
                                "fit": "crop",
                                "caption": "血球形態影像",
                                "sourceLabel": "院內教學素材",
                                "altText": "血球形態",
                            }
                        ],
                    }
                ],
                output_path=output,
                provenance={"sourceMaterialId": "mat-1", "sourceDraftId": "draft-1"},
                image_resolver=lambda _asset: source,
                quality_report=report,
            )

            self.assertTrue((root / "source-crop.png").is_file())
            rendered = runtime.Presentation(str(output))
            self.assertEqual(len(rendered.slides), 2)
            slide = rendered.slides[1]
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            text = _slide_text(slide)
            self.assertIn("血球形態影像｜院內教學素材", text)
            self.assertNotIn("MISSING_IMAGE", {item["code"] for item in report["warnings"]})
            self.assertEqual(report["slideCount"], 2)
            self.assertEqual(report["errorCount"], 0)


if __name__ == "__main__":
    unittest.main()
