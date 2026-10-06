"""Validate downloaded artifacts, not API success indicators."""
import json
from pathlib import Path
import sys
import wave
import zipfile

path = Path(sys.argv[1])
kind = sys.argv[2]
assert path.stat().st_size > 1000
if kind == "pptx":
    from pptx import Presentation
    deck = Presentation(path)
    assert len(deck.slides) >= 2
    notes = [slide.notes_slide.notes_text_frame.text for slide in deck.slides]
    assert any(text.strip() for text in notes)
    if len(sys.argv) > 3:
        expected = sys.argv[3]
        assert any(expected in text for text in notes), "Saved script missing from actual PPTX speaker notes"
        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
        from teacher_app.materials.ai_video_runtime import narration_for_slide
        assert narration_for_slide({"speakerNotes": expected}) == " ".join(expected.split())
    with zipfile.ZipFile(path) as package:
        assert any("notesSlides/notesSlide" in name for name in package.namelist())
    print(json.dumps({"slides": len(deck.slides), "notes": notes, "bytes": path.stat().st_size}))
elif kind == "wav":
    with wave.open(str(path)) as audio:
        assert audio.getnframes() > 0
        assert audio.getframerate() > 0
        print(json.dumps({"duration": audio.getnframes()/audio.getframerate(), "bytes": path.stat().st_size}))
else:
    raise ValueError(kind)
