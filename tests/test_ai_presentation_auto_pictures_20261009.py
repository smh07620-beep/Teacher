"""Slides pick up pictures from the teacher's own sources (Word / PowerPoint / PDF / image files)."""
import io
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from PIL import Image

from teacher_app.materials import ai_presentation_images as pics
from teacher_app.materials import ai_presentation_runtime as runtime
from teacher_app.materials.ai_presentation_storage import PresentationStorage

try:
    import pymupdf
except ImportError:  # pragma: no cover
    pymupdf = None

needs_pptx = unittest.skipIf(runtime.Presentation is None, "python-pptx is installed only on the AI Worker")
needs_pdf = unittest.skipIf(pymupdf is None, "PyMuPDF not installed")


def _png_bytes(color, size=(320, 240)) -> bytes:
    image = Image.new("RGB", size, color)
    for x in range(0, size[0], 8):  # a little structure so files are not identical across colors
        for y in range(0, size[1], 40):
            image.putpixel((x, y), (255 - color[0], 255 - color[1], color[2]))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


SMEAR = "推片角度約三十到四十五度，抹片製作要薄而均勻"
URINE = "尿液沉渣的草酸鈣結晶與尿酸結晶判讀"


def _docx(path: Path, sections: list[tuple[str, bytes]]) -> Path:
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')
    body, rels = [], []
    for index, (text, _data) in enumerate(sections, 1):
        rid = f"rId{index + 10}"
        body.append(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>")
        body.append(f'<w:p><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><a:blip r:embed="{rid}"/>'
                    "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>")
        rels.append(f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                    f'Target="media/image{index}.png"/>')
    document = f'<?xml version="1.0" encoding="UTF-8"?><w:document {ns}><w:body>{"".join(body)}</w:body></w:document>'
    rels_xml = ('<?xml version="1.0" encoding="UTF-8"?><Relationships '
                'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + "".join(rels) + "</Relationships>")
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", rels_xml)
        for index, (_text, data) in enumerate(sections, 1):
            archive.writestr(f"word/media/image{index}.png", data)
    return path


def _pptx(path: Path, pages: list[tuple[str, bytes]], *, logo: bytes | None = None) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for text, data in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = text
        stream = io.BytesIO(data)
        slide.shapes.add_picture(stream, Inches(1), Inches(2), Inches(4), Inches(3))
        if logo:
            slide.shapes.add_picture(io.BytesIO(logo), Inches(8), Inches(0.1), Inches(1), Inches(1))
    prs.save(str(path))
    return path


def _pdf(path: Path, pages: list[tuple[str, bytes]]) -> Path:
    document = pymupdf.open()
    for text, data in pages:
        page = document.new_page()
        page.insert_text((72, 72), text, fontname="china-t")
        page.insert_image(pymupdf.Rect(72, 120, 372, 300), stream=data)
    document.save(str(path))
    document.close()
    return path


def _slide(slide_id, title, bullets, layout="content", blocks=None):
    return {"id": slide_id, "order": 1, "enabled": True, "title": title, "bullets": bullets,
            "layout": layout, "blocks": blocks or [], "speakerNotes": ""}


class NormalizeImageTests(unittest.TestCase):
    def test_small_banner_and_normal_images(self):
        self.assertIsNone(pics._normalize_image(_png_bytes((10, 20, 30), size=(48, 48))))  # icon
        self.assertIsNone(pics._normalize_image(_png_bytes((10, 20, 30), size=(1200, 120))))  # divider banner
        ok = pics._normalize_image(_png_bytes((10, 20, 30)))
        self.assertIsNotNone(ok)
        self.assertEqual(ok[1], "image/png")

    def test_large_picture_is_scaled_down_and_transparency_flattened(self):
        big = io.BytesIO()
        Image.new("RGBA", (3200, 2400), (0, 0, 0, 0)).save(big, "PNG")
        data, mime, width, height = pics._normalize_image(big.getvalue())
        self.assertLessEqual(max(width, height), pics.MAX_EDGE)
        self.assertIn(mime, {"image/png", "image/jpeg"})
        with Image.open(io.BytesIO(data)) as image:
            self.assertEqual(image.mode in {"RGB", "L"}, True)

    def test_garbage_bytes_are_ignored(self):
        self.assertIsNone(pics._normalize_image(b"not an image"))


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def test_docx_images_keep_surrounding_text(self):
        path = _docx(self.tmp / "a.docx", [(SMEAR, _png_bytes((200, 40, 40))), (URINE, _png_bytes((40, 200, 40)))])
        rows = pics.extract_from_file(path, {})
        self.assertEqual(len(rows), 2)
        self.assertIn("推片", rows[0][1])
        self.assertIn("尿液", rows[1][1])

    @needs_pptx
    def test_pptx_pictures_with_slide_text_and_repeated_logo_skipped(self):
        logo = _png_bytes((1, 2, 3), size=(300, 300))
        path = _pptx(self.tmp / "a.pptx", [(SMEAR, _png_bytes((200, 40, 40))), (URINE, _png_bytes((40, 200, 40)))], logo=logo)
        rows = pics.extract_from_file(path, {})
        self.assertEqual(len(rows), 2, "the logo repeated on every slide must not become a candidate")
        self.assertIn("推片", rows[0][1])

    @needs_pdf
    def test_pdf_pictures_with_page_text(self):
        path = _pdf(self.tmp / "a.pdf", [(SMEAR, _png_bytes((200, 40, 40))), (URINE, _png_bytes((40, 200, 40)))])
        rows = pics.extract_from_file(path, {})
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row[0][:4] == b"\x89PNG" for row in rows))

    def test_plain_image_file_uses_its_title_as_context(self):
        path = self.tmp / "figure.png"
        path.write_bytes(_png_bytes((90, 90, 200)))
        rows = pics.extract_from_file(path, {"title": "尿液沉渣結晶圖", "filename": "figure.png"})
        self.assertEqual(len(rows), 1)
        self.assertIn("尿液沉渣", rows[0][1])

    def test_unsupported_types_give_nothing(self):
        path = self.tmp / "data.xlsx"
        path.write_bytes(b"x")
        self.assertEqual(pics.extract_from_file(path, {}), [])


class MatchTests(unittest.TestCase):
    def _candidates(self):
        return [
            {"context": URINE, "path": Path("u"), "sha256": "u" * 64, "mimeType": "image/png", "caption": "", "label": "A", "materialId": "a"},
            {"context": SMEAR, "path": Path("s"), "sha256": "s" * 64, "mimeType": "image/png", "caption": "", "label": "A", "materialId": "a"},
        ]

    def test_each_picture_goes_to_its_own_slide(self):
        slides = [_slide("s1", "抹片製作", ["推片角度三十到四十五度", "抹片要薄而均勻"]),
                  _slide("s2", "尿液沉渣", ["草酸鈣結晶", "尿酸結晶判讀"])]
        got = pics.match_pictures(slides, self._candidates())
        self.assertEqual([(s, c) for s, c, _ in got], [(0, 1), (1, 0)])

    def test_unrelated_picture_is_left_out(self):
        slides = [_slide("s1", "細菌培養", ["血液培養瓶採檢", "革蘭氏染色"])]
        self.assertEqual(pics.match_pictures(slides, self._candidates()), [])

    def test_realistic_lecture_matches_each_figure_and_skips_the_off_topic_one(self):
        slides = [_slide("1", "血液抹片的製作", ["推片角度約 30–45 度", "血滴大小與推片速度影響厚薄", "抹片須自然風乾後再染色"]),
                  _slide("2", "瑞氏染色原理", ["瑞氏染劑含酸性與鹼性染料", "緩衝液 pH 6.8 影響染色結果"]),
                  _slide("3", "白血球分類", ["中性球、淋巴球、單核球", "嗜酸性球與嗜鹼性球"]),
                  _slide("4", "尿液沉渣檢查", ["離心後取沉渣鏡檢", "常見結晶：草酸鈣、尿酸"])]
        contexts = [
            "第三章 抹片製作 載玻片置於平台，血滴置於一端，以另一玻片呈 30 至 45 度角向後接觸血滴後平穩推出，形成舌狀薄層。圖 3-1 推片示意圖",
            "染色時先滴上瑞氏染劑固定 1 分鐘，再加入緩衝液混合 5 分鐘後沖洗。圖 3-2 瑞氏染色後的血液抹片",
            "圖 4-1 周邊血液中性球形態，細胞核分葉為二至五葉，細胞質呈淡粉色，內含細小顆粒。淋巴球較小，核圓而深染。",
            "圖 5-2 尿沉渣中的草酸鈣結晶，呈信封狀，常見於酸性尿。離心 5 分鐘後取沉渣一滴置於載玻片鏡檢。",
            "圖 9 公司組織圖與院內分機表 實驗室主管 品管 行政 電話 內線",
        ]
        candidates = [{"context": text} for text in contexts]
        got = pics.match_pictures(slides, candidates)
        self.assertEqual([(s, c) for s, c, _ in got], [(0, 0), (1, 1), (2, 2), (3, 3)])

    def test_picture_that_fits_two_slides_equally_is_not_guessed(self):
        slides = [_slide("1", "抹片製作", ["推片角度與抹片製作"]), _slide("2", "抹片製作", ["推片角度與抹片製作"])]
        self.assertEqual(pics.match_pictures(slides, [{"context": SMEAR}]), [])

    def test_only_plain_text_slides_are_eligible(self):
        busy = _slide("s1", "抹片製作", ["推片角度"], layout="table", blocks=[{"type": "table", "headers": ["a"], "rows": []}])
        has_image = _slide("s2", "抹片製作", ["推片角度"], blocks=[{"type": "image"}])
        cover = _slide("s3", "抹片製作", [], layout="title")
        long_text = _slide("s4", "抹片製作", ["推片角度" * 120])
        disabled = dict(_slide("s5", "抹片製作", ["推片角度"]), enabled=False)
        for slide in (busy, has_image, cover, long_text, disabled):
            self.assertFalse(pics.slide_accepts_picture(slide), slide["id"])
        self.assertTrue(pics.slide_accepts_picture(_slide("ok", "抹片製作", ["推片角度"])))


@needs_pdf
class AttachTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.docx = _docx(self.tmp / "main.docx", [(SMEAR, _png_bytes((200, 40, 40)))])
        self.pdf = _pdf(self.tmp / "ref.pdf", [(URINE, _png_bytes((40, 200, 40)))])
        self.stored = []

    def _fetch(self, material):
        if material["id"] == "broken":
            raise RuntimeError("storage offline")
        root = Path(tempfile.mkdtemp(dir=self.tmp))
        target = root / ("source" + Path(material["filename"]).suffix)
        shutil.copy2({"main": self.docx, "ref": self.pdf}[material["id"]], target)
        return root, target

    def _store(self, path, sha, mime):
        self.stored.append(sha)
        return {"backend": "r2", "key": f"ai-presentations/images/{sha}.png", "sha256": sha, "mimeType": mime}

    def test_pictures_from_main_and_reference_sources_reach_matching_slides(self):
        slides = [_slide("s1", "抹片製作", ["推片角度三十到四十五度"]), _slide("s2", "尿液沉渣", ["草酸鈣結晶與尿酸結晶"]),
                  _slide("s3", "總結", ["重點回顧"], layout="summary")]
        materials = [{"id": "main", "title": "血液學講義", "filename": "main.docx"},
                     {"id": "ref", "title": "尿液簡報", "filename": "ref.pdf"},
                     {"id": "broken", "title": "壞檔", "filename": "x.docx"}]
        result, info = pics.attach_pictures(slides, materials, workdir=self.tmp / "work", fetch_source=self._fetch, store_image=self._store)
        self.assertEqual(info["applied"], 2)
        self.assertEqual([b["sourceLabel"] for s in result for b in s["blocks"]], ["血液學講義", "尿液簡報"])
        self.assertEqual(result[2]["blocks"], [], "summary slide stays text only")
        self.assertEqual([item["label"] for item in info["skipped"]], ["壞檔"], "one broken source is reported, not fatal")
        for slide in result[:2]:
            asset = slide["blocks"][0]["asset"]
            self.assertTrue(runtime._IMAGE_KEY.fullmatch(asset["key"]))
        # blocks survive the normal slide validation, so revisions keep the pictures
        normalized = runtime.normalize_slides(result)
        self.assertTrue(all(slide["blocks"][0].get("asset") for slide in normalized[:2]))

    def test_store_failure_skips_only_that_picture(self):
        slides = [_slide("s1", "抹片製作", ["推片角度三十到四十五度"])]

        def fail(path, sha, mime):
            raise RuntimeError("bucket full")

        result, info = pics.attach_pictures(slides, [{"id": "main", "title": "講義", "filename": "main.docx"}],
                                            workdir=self.tmp / "work", fetch_source=self._fetch, store_image=fail)
        self.assertEqual(info["applied"], 0)
        self.assertEqual(result[0]["blocks"], [])

    def test_no_work_when_no_slide_can_take_a_picture(self):
        fetch = MagicMock()
        _, info = pics.attach_pictures([_slide("s1", "封面", [], layout="title")], [{"id": "main", "filename": "main.docx"}],
                                       workdir=self.tmp / "work", fetch_source=fetch, store_image=self._store)
        fetch.assert_not_called()
        self.assertEqual(info["applied"], 0)


@needs_pptx
class RenderLayoutTests(unittest.TestCase):
    def test_bullets_stay_left_and_picture_goes_right(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "p.png"
            image.write_bytes(_png_bytes((20, 120, 200)))
            output = root / "out.pptx"
            slide = _slide("s1", "抹片製作", ["推片角度三十到四十五度", "抹片要薄而均勻"], blocks=[{
                "type": "image", "asset": {"backend": "r2", "key": "ai-presentations/images/" + "a" * 64 + ".png",
                                           "sha256": "a" * 64, "mimeType": "image/png"},
                "caption": "抹片示意", "sourceLabel": "血液學講義", "fit": "contain", "altText": "抹片"}])
            runtime.render_pptx(title="t", slides=[slide], output_path=output,
                                provenance={"sourceMaterialId": "m", "sourceDraftId": "d"}, image_resolver=lambda _a: image)
            deck = runtime.Presentation(str(output))
            page = deck.slides[1]
            pictures = [s for s in page.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            text_boxes = [s for s in page.shapes if s.has_text_frame and "推片角度" in s.text_frame.text]
            self.assertEqual(len(text_boxes), 1)
            self.assertLessEqual(text_boxes[0].left + text_boxes[0].width, pictures[0].left, "picture must not cover the bullets")
            caption = [s for s in page.shapes if s.has_text_frame and "血液學講義" in s.text_frame.text]
            self.assertTrue(caption, "the source file title is shown under the picture")


class GeneratePresentationPicturesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.docx = _docx(self.tmp / "main.docx", [(SMEAR, _png_bytes((200, 40, 40)))])

    def _run(self, job_request=None, backend="r2"):
        draft = {"id": "d1", "draftType": "slides", "status": "approved", "approvedBy": "t", "approvedAt": "2026-10-09",
                 "group": "g", "area": "a", "materialId": "m1", "title": "血液抹片", "sourceJobId": "sj1",
                 "body": "第 1 張：抹片製作\n- 推片角度三十到四十五度\n- 抹片要薄而均勻\n第 2 張：總結\n- 重點回顧"}
        source = {"id": "m1", "group": "g", "area": "a", "title": "血液學講義", "filename": "main.docx"}
        captured = {}
        storage = MagicMock()
        storage.image_backend.return_value = backend
        storage.store.return_value = {"backend": "r2", "key": "k", "filename": "f.pptx", "sha256": "x", "byteSize": 1, "mimeType": "m"}
        storage.store_image.side_effect = lambda path, sha256, mime_type: {
            "backend": "r2", "key": f"ai-presentations/images/{sha256}.png", "sha256": sha256, "mimeType": mime_type}

        def fake_render(**kwargs):
            captured["rendered"] = kwargs["slides"]
            kwargs["quality_report"].update({"warnings": [], "errors": [], "slideCount": len(kwargs["slides"])})
            kwargs["output_path"].write_bytes(b"x")

        def fetch(entry, **_kw):
            root = Path(tempfile.mkdtemp(dir=self.tmp))
            target = root / "source.docx"
            shutil.copy2(self.docx, target)
            return root, target

        with patch.object(runtime.repository, "get_presentation_by_source_job_id", return_value=None), \
             patch.object(runtime.media_script_repository, "get_script", return_value=draft), \
             patch.object(runtime.media_script_repository, "list_scripts", return_value=[]), \
             patch.object(runtime.media_script_repository, "get_job", return_value={"request": {"referenceMaterialIds": []}}), \
             patch.object(runtime.material_repository, "get_material", return_value=source), \
             patch.object(runtime, "_validated_template", return_value=None), \
             patch("teacher_app.assessments.ai_runtime.material_source_to_temp", side_effect=fetch), \
             patch.object(runtime, "render_pptx", side_effect=fake_render), \
             patch.object(runtime.repository, "create_presentation", side_effect=lambda **kw: captured.update(saved=kw) or {"id": "p1"}):
            runtime.generate_presentation(job={"id": "j1", "draftId": "d1", "group": "g", "area": "a", "request": job_request or {}}, storage=storage)
        return captured

    def test_matching_slide_gets_picture_saved_with_the_revision_and_review_warning(self):
        got = self._run()
        first, second = got["rendered"][0], got["rendered"][1]
        self.assertEqual(len(first.get("blocks") or []), 1)
        self.assertEqual(first["blocks"][0]["sourceLabel"], "血液學講義")
        self.assertFalse(second.get("blocks"))
        self.assertEqual(got["saved"]["slides"][0]["blocks"], first["blocks"], "revision keeps the picture for later edits")
        codes = [item["code"] for item in got["saved"]["quality_manifest"]["warnings"]]
        self.assertIn("AUTO_PICTURES", codes)

    def test_teacher_can_switch_pictures_off_per_job(self):
        got = self._run({"autoPictures": False})
        self.assertTrue(all(not slide.get("blocks") for slide in got["rendered"]))

    def test_environment_switch_turns_pictures_off(self):
        with patch.dict("os.environ", {"AI_PRESENTATION_AUTO_PICTURES": "false"}):
            got = self._run()
        self.assertTrue(all(not slide.get("blocks") for slide in got["rendered"]))

    def test_provider_without_object_keys_skips_pictures_but_still_renders(self):
        got = self._run(backend="")
        self.assertEqual(len(got["rendered"]), 2)
        self.assertTrue(all(not slide.get("blocks") for slide in got["rendered"]))


class StorageImageTests(unittest.TestCase):
    def test_image_is_stored_content_addressed_on_object_storage(self):
        storage = PresentationStorage(storage_adapter=MagicMock())
        client = MagicMock()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "p.png"
            path.write_bytes(_png_bytes((1, 2, 3)))
            with patch.object(PresentationStorage, "backend", return_value="r2"), \
                 patch("teacher_app.materials.ai_presentation_storage.providers.r2_client", return_value=client), \
                 patch("teacher_app.materials.ai_presentation_storage.r2_ledger.record_object") as ledger:
                asset = storage.store_image(path, sha256="b" * 64, mime_type="image/png")
        self.assertEqual(asset["key"], "ai-presentations/images/" + "b" * 64 + ".png")
        self.assertEqual(client.upload_file.call_args.kwargs["ExtraArgs"], {"ContentType": "image/png"})
        ledger.assert_called_once()
        self.assertTrue(runtime._IMAGE_KEY.fullmatch(asset["key"]))

    def test_providers_without_object_keys_do_not_accept_pictures(self):
        storage = PresentationStorage(storage_adapter=MagicMock())
        with patch.object(PresentationStorage, "backend", return_value="mega"):
            self.assertEqual(storage.image_backend(), "")
            with self.assertRaises(RuntimeError):
                storage.store_image(Path("x.png"), sha256="c" * 64, mime_type="image/png")


class RetentionPictureTests(unittest.TestCase):
    def test_picture_assets_are_read_from_slides_json(self):
        from teacher_app.maintenance import retention

        key = "ai-presentations/images/" + "d" * 64 + ".jpg"
        slides = [{"blocks": [{"type": "image", "asset": {"backend": "r2", "key": key}},
                              {"type": "table"}, {"type": "image", "asset": {"backend": "r2", "key": "elsewhere/x.png"}}]}]
        import json
        self.assertEqual(retention._picture_assets(json.dumps(slides)), {("r2", key)})
        self.assertEqual(retention._picture_assets(slides), {("r2", key)})
        self.assertEqual(retention._picture_assets("not json"), set())


if __name__ == "__main__":
    unittest.main()
