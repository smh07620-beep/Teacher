"""Valid tiny source documents for browser upload, no generated results seeded."""
def build(root):
    import fitz
    from docx import Document
    from pptx import Presentation
    from PIL import Image, ImageDraw
    text = "E2E laboratory specimen identity safety quality control and sample collection time. " * 5
    (root / "pasted.txt").write_text(text, encoding="utf-8")
    pdf = fitz.open(); page = pdf.new_page(); page.insert_textbox(fitz.Rect(50,50,500,750),text)
    pdf.save(root / "source.pdf"); pdf.close()
    word = Document(); word.add_paragraph(text); word.save(root / "source.docx")
    ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[1]); slide.shapes.title.text = "E2E source"
    slide.placeholders[1].text = text; slide.notes_slide.notes_text_frame.text = text; ppt.save(root / "source.pptx")
    image = Image.new("RGB",(640,360),"white"); ImageDraw.Draw(image).text((20,20),"E2E specimen safety identity",fill="black")
    image.save(root / "source.png")
