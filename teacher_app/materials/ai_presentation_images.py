"""Automatically pick pictures for AI PowerPoint slides from the teacher's own sources.

Teachers should not have to hunt for figures. While a deck is generated we look
inside the material the deck was written from, plus the reference files the
teacher added (Word, PowerPoint, PDF or plain image files), collect their
pictures together with the text printed around each picture, and give every
slide the single best-matching picture.

Matching is local and deterministic (character-bigram similarity between a
picture's surrounding text and the slide's text); pictures are never sent to an
AI provider. Slides stay drafts until the teacher reviews them, and every
picture carries its source file title so the teacher can see where it came from.
"""
from __future__ import annotations

import hashlib
import io
import logging
import shutil
import zipfile
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from teacher_app.materials import script_alignment

LOGGER = logging.getLogger(__name__)

IMAGE_FILE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
MIN_SIDE = 160            # px; smaller pictures are icons / bullets / logos
MIN_AREA = 48_000         # px²
MAX_ASPECT = 5.0          # banners and divider lines are not teaching figures
MAX_EDGE = 1600           # px; larger pictures are scaled down before storing
MAX_BYTES = 3 * 1024 * 1024
MAX_PER_SOURCE = 60
MAX_PAGES_SCANNED = 120
MAX_CANDIDATES = 160
MIN_MATCH_SCORE = 0.10    # below this a picture is left out rather than guessed
AMBIGUITY_RATIO = 1.25    # a picture must fit its slide clearly better than any other slide
MAX_SLIDE_BULLETS = 5     # picture sits beside the text, so keep the text short
MAX_SLIDE_CHARS = 330
CONTEXT_LIMIT = 700
_ELIGIBLE_LAYOUTS = {"content", "image"}
_BLOCKING_BLOCKS = {"image", "table", "chart", "comparison"}


def _normalize_image(data: bytes) -> tuple[bytes, str, int, int] | None:
    """Return (bytes, mime, width, height) for a usable picture, else None."""
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover - Pillow ships with the app
        return None
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            width, height = image.size
            if min(width, height) < MIN_SIDE or width * height < MIN_AREA:
                return None
            if max(width / height, height / width) > MAX_ASPECT:
                return None
            has_alpha = image.mode in {"RGBA", "LA"} or (image.mode == "P" and "transparency" in image.info)
            if has_alpha:
                background = Image.new("RGB", image.size, (255, 255, 255))
                background.paste(image.convert("RGBA"), mask=image.convert("RGBA").split()[-1])
                image = background
            elif image.mode not in {"RGB", "L"}:
                image = image.convert("RGB")
            if max(image.size) > MAX_EDGE:
                scale = MAX_EDGE / float(max(image.size))
                image = image.resize((max(1, round(image.size[0] * scale)), max(1, round(image.size[1] * scale))), Image.LANCZOS)
            buffer = io.BytesIO()
            as_png = str(getattr(image, "format", "") or "").upper() == "PNG" or data[:8] == b"\x89PNG\r\n\x1a\n"
            if as_png:
                image.save(buffer, format="PNG", optimize=True)
                if buffer.tell() > MAX_BYTES // 2:  # photos are far smaller as JPEG
                    buffer = io.BytesIO()
                    image.convert("RGB").save(buffer, format="JPEG", quality=88, optimize=True)
                    mime = "image/jpeg"
                else:
                    mime = "image/png"
            else:
                image.convert("RGB").save(buffer, format="JPEG", quality=88, optimize=True)
                mime = "image/jpeg"
            payload = buffer.getvalue()
            if not payload or len(payload) > MAX_BYTES:
                return None
            return payload, mime, image.size[0], image.size[1]
    except Exception:
        return None


def _clean_context(text: str) -> str:
    return " ".join(str(text or "").replace("\x00", " ").split())[:CONTEXT_LIMIT]


# ── extractors: each yields (raw image bytes, surrounding text, short caption) ──

def _from_docx(path: Path) -> list[tuple[bytes, str, str]]:
    from teacher_app.atlas.importer import _docx_inline_images

    found = _docx_inline_images(path)
    rows: list[tuple[bytes, str, str]] = []
    with zipfile.ZipFile(path) as archive:
        for item in found[:MAX_PER_SOURCE]:
            try:
                rows.append((archive.read(item["mediaPath"]), item.get("section") or "", item.get("caption") or ""))
            except KeyError:
                continue
    return rows


def _shape_pictures(shapes) -> list[bytes]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    blobs: list[bytes] = []
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                blobs.extend(_shape_pictures(shape.shapes))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blobs.append(shape.image.blob)
        except Exception:
            continue
    return blobs


def _shape_text(shapes) -> str:
    parts: list[str] = []
    for shape in shapes:
        try:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if hasattr(shape, "shapes"):
                parts.append(_shape_text(shape.shapes))
        except Exception:
            continue
    return " ".join(parts)


def _from_pptx(path: Path) -> list[tuple[bytes, str, str]]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    per_slide: list[tuple[list[bytes], str, str]] = []
    seen: dict[str, int] = {}
    for slide in presentation.slides:
        blobs = _shape_pictures(slide.shapes)
        text = _shape_text(slide.shapes)
        title = ""
        try:
            title = slide.shapes.title.text if slide.shapes.title is not None else ""
        except Exception:
            title = ""
        for blob in blobs:
            digest = hashlib.sha256(blob).hexdigest()
            seen[digest] = seen.get(digest, 0) + 1
        per_slide.append((blobs, text, title))
    rows: list[tuple[bytes, str, str]] = []
    used: set[str] = set()
    for blobs, text, title in per_slide:
        for blob in blobs:
            digest = hashlib.sha256(blob).hexdigest()
            # A picture repeated on several slides is a logo / background.
            if seen.get(digest, 0) > 1 or digest in used:
                continue
            used.add(digest)
            rows.append((blob, text, title))
            if len(rows) >= MAX_PER_SOURCE:
                return rows
    return rows


def _from_pdf(path: Path) -> list[tuple[bytes, str, str]]:
    import pymupdf

    rows: list[tuple[bytes, str, str]] = []
    xref_pages: dict[int, int] = {}
    with pymupdf.open(str(path)) as document:
        pages = min(len(document), MAX_PAGES_SCANNED)
        page_images: list[tuple[list[int], str]] = []
        for index in range(pages):
            page = document[index]
            xrefs = [int(item[0]) for item in page.get_images(full=True)]
            for xref in xrefs:
                xref_pages[xref] = xref_pages.get(xref, 0) + 1
            page_images.append((xrefs, page.get_text("text") or ""))
        done: set[int] = set()
        for xrefs, text in page_images:
            for xref in xrefs:
                if xref in done or xref_pages.get(xref, 0) > 1:
                    continue  # repeated on many pages: logo / watermark
                done.add(xref)
                try:
                    pixmap = pymupdf.Pixmap(document, xref)
                    if pixmap.colorspace is not None and pixmap.n - pixmap.alpha >= 4:
                        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pixmap)
                    if pixmap.alpha:
                        pixmap = pymupdf.Pixmap(pixmap, 0)
                    rows.append((pixmap.tobytes("png"), text, ""))
                except Exception:
                    continue
                if len(rows) >= MAX_PER_SOURCE:
                    return rows
    return rows


def _from_image_file(path: Path, material: Mapping[str, Any]) -> list[tuple[bytes, str, str]]:
    label = " ".join(str(material.get(key) or "") for key in ("title", "filename", "description"))
    return [(path.read_bytes(), label, str(material.get("title") or "")[:120])]


def extract_from_file(path: Path, material: Mapping[str, Any]) -> list[tuple[bytes, str, str]]:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return _from_docx(path)
    if suffix == ".pptx":
        return _from_pptx(path)
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix in IMAGE_FILE_EXTENSIONS:
        return _from_image_file(path, material)
    return []


def material_label(material: Mapping[str, Any]) -> str:
    return " ".join(str(material.get("title") or material.get("filename") or "教材").split())[:80]


def collect_candidates(
    materials: Sequence[Mapping[str, Any]],
    workdir: Path,
    *,
    fetch_source: Callable[[Mapping[str, Any]], tuple[Path, Path]],
    max_total: int = MAX_CANDIDATES,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    """Return (candidates, warnings). One unreadable source never blocks the others."""
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    candidates: list[dict[str, Any]] = []
    warnings: list[dict[str, str]] = []
    seen: set[str] = set()
    for material in materials:
        if len(candidates) >= max_total:
            break
        extension = Path(str(material.get("filename") or material.get("storageFilename") or "")).suffix.lower()
        if extension not in {".docx", ".pptx", ".pdf", *IMAGE_FILE_EXTENSIONS}:
            continue
        temp_root: Path | None = None
        try:
            temp_root, source = fetch_source(material)
            for raw, context, caption in extract_from_file(Path(source), material):
                normalized = _normalize_image(raw)
                if normalized is None:
                    continue
                payload, mime, width, height = normalized
                digest = hashlib.sha256(payload).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                target = workdir / f"cand-{len(candidates) + 1}.{'png' if mime == 'image/png' else 'jpg'}"
                target.write_bytes(payload)
                candidates.append({
                    "path": target, "sha256": digest, "mimeType": mime, "width": width, "height": height,
                    "context": _clean_context(f"{caption} {context}"), "caption": _clean_context(caption)[:120],
                    "materialId": str(material.get("id") or ""), "label": material_label(material),
                })
                if len(candidates) >= max_total:
                    break
        except Exception as exc:
            LOGGER.warning("slide picture scan skipped material=%s error=%s", str(material.get("id") or "")[:60], type(exc).__name__)
            warnings.append({"materialId": str(material.get("id") or ""), "label": material_label(material), "reason": type(exc).__name__})
        finally:
            if temp_root:
                shutil.rmtree(temp_root, ignore_errors=True)
    return candidates, warnings


# ── matching ──

def slide_accepts_picture(slide: Mapping[str, Any]) -> bool:
    """Only plain text slides get a picture; tables/charts/comparisons already own the space."""
    if not bool(slide.get("enabled", True)):
        return False
    if str(slide.get("layout") or "content").lower() not in _ELIGIBLE_LAYOUTS:
        return False
    if any(isinstance(block, Mapping) and str(block.get("type") or "").lower() in _BLOCKING_BLOCKS for block in list(slide.get("blocks") or [])):
        return False
    bullets = [str(item or "") for item in list(slide.get("bullets") or []) if str(item or "").strip()]
    if not bullets and not str(slide.get("title") or "").strip():
        return False
    return len(bullets) <= MAX_SLIDE_BULLETS and sum(len(item) for item in bullets) <= MAX_SLIDE_CHARS


def match_pictures(
    slides: Sequence[Mapping[str, Any]],
    candidates: Sequence[Mapping[str, Any]],
    *,
    min_score: float = MIN_MATCH_SCORE,
) -> list[tuple[int, int, float]]:
    """Greedy one-to-one (slide index, candidate index, score) assignment, best scores first."""
    eligible = [index for index, slide in enumerate(slides) if slide_accepts_picture(slide)]
    if not eligible or not candidates:
        return []
    matrix = script_alignment._similarity_matrix([str(item.get("context") or "") for item in candidates], list(slides))
    pairs = []
    for c in range(len(candidates)):
        for s in eligible:
            score = matrix[c][s]
            rival = max((matrix[c][other] for other in eligible if other != s), default=0.0)
            # Unclear which slide a picture belongs to -> leave it out rather than guess.
            if score >= min_score and score >= AMBIGUITY_RATIO * rival:
                pairs.append((score, s, c))
    pairs.sort(key=lambda pair: (-pair[0], pair[1], pair[2]))
    used_slides: set[int] = set()
    used_candidates: set[int] = set()
    chosen: list[tuple[int, int, float]] = []
    for score, slide_index, candidate_index in pairs:
        if slide_index in used_slides or candidate_index in used_candidates:
            continue
        used_slides.add(slide_index)
        used_candidates.add(candidate_index)
        chosen.append((slide_index, candidate_index, round(float(score), 4)))
    return sorted(chosen)


def attach_pictures(
    slides: Sequence[Mapping[str, Any]],
    materials: Sequence[Mapping[str, Any]],
    *,
    workdir: Path,
    fetch_source: Callable[[Mapping[str, Any]], tuple[Path, Path]],
    store_image: Callable[[Path, str, str], dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Return (slides with picture blocks added, summary). Never raises for bad sources."""
    result = [dict(slide) for slide in slides]
    summary: dict[str, Any] = {"applied": 0, "candidateCount": 0, "sourceCount": len(materials), "skipped": [], "attached": []}
    if not any(slide_accepts_picture(slide) for slide in result):
        return result, summary
    candidates, warnings = collect_candidates(materials, workdir, fetch_source=fetch_source)
    summary["candidateCount"] = len(candidates)
    summary["skipped"] = warnings
    for slide_index, candidate_index, score in match_pictures(result, candidates):
        candidate = candidates[candidate_index]
        try:
            asset = store_image(Path(candidate["path"]), candidate["sha256"], candidate["mimeType"])
        except Exception as exc:
            LOGGER.warning("slide picture not stored error=%s", type(exc).__name__)
            summary["skipped"].append({"materialId": candidate["materialId"], "label": candidate["label"], "reason": type(exc).__name__})
            continue
        slide = dict(result[slide_index])
        blocks = list(slide.get("blocks") or [])
        blocks.append({
            "type": "image", "asset": asset, "fit": "contain",
            "caption": candidate["caption"], "altText": str(slide.get("title") or "")[:200],
            "sourceLabel": candidate["label"], "title": "",
        })
        slide["blocks"] = blocks
        result[slide_index] = slide
        summary["applied"] += 1
        summary["attached"].append({"slideId": str(slide.get("id") or ""), "materialId": candidate["materialId"], "score": score})
    return result, summary


__all__ = [
    "attach_pictures", "collect_candidates", "extract_from_file", "match_pictures",
    "slide_accepts_picture", "material_label", "MIN_MATCH_SCORE",
]
