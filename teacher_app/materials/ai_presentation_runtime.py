"""AI PowerPoint renderer and worker-side generation orchestration."""
from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import tempfile
import time
from pathlib import Path
from typing import Any

from teacher_app.materials import ai_presentation_images as picture_picker
from teacher_app.materials import ai_presentation_repository as repository
from teacher_app.materials import ai_presentation_quality as quality
from teacher_app.materials.ai_presentation_storage import PresentationStorage
from teacher_app.materials import media_script_repository
from teacher_app.materials import script_alignment
from teacher_app.materials import repository as material_repository

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
except ImportError:  # pragma: no cover
    Presentation = None
    Inches = Pt = MSO_SHAPE = CategoryChartData = XL_CHART_TYPE = None

try:  # Pillow is already optional in other material-image paths.
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

LOGGER = logging.getLogger(__name__)
_MAX_SLIDES = 60
_MAX_BULLETS = 120
_SECRET_MARKERS = (
    "token=", "password=", "secret=", "api_key=", "apikey=", "authorization:",
    "bearer ", "database_url=", "r2_secret", "mega_password", "client_secret=",
)
_LOCAL_PATH_PATTERN = re.compile(r"(?:[a-z]:[\\/]|/(?:home|tmp|var|mnt|etc)/|file://)", re.I)
_LAYOUT_KINDS = {"title", "section", "content", "image", "comparison", "table", "summary"}
_BLOCK_KINDS = {"image", "chart", "table", "comparison", "callout"}
_IMAGE_MIME_TYPES = {"image/png", "image/jpeg"}
_IMAGE_KEY = re.compile(r"^ai-presentations/images/[a-z0-9][a-z0-9._/-]{0,220}$", re.I)


def _clean(value: Any, limit: int) -> str:
    text = str(value or "").replace("\x00", " ").strip()
    text = "\n".join(" ".join(line.split()) for line in text.splitlines())
    return text[:limit]


def _sensitive_provenance_value(value: Any) -> bool:
    if isinstance(value, list):
        return any(_sensitive_provenance_value(item) for item in value)
    text = str(value or "")
    lowered = text.lower()
    return any(marker in lowered for marker in _SECRET_MARKERS) or bool(_LOCAL_PATH_PATTERN.search(text))


def _provenance(payload: dict[str, Any]) -> str:
    return json.dumps(repository.sanitize_provenance(payload), ensure_ascii=False, separators=(",", ":"))


def _core_provenance_comment(payload: dict[str, Any]) -> str:
    """Return a python-pptx-safe locator while the full provenance stays in notes/DB."""
    full = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if len(full) <= 255:
        return full
    compact = {
        "sourceMaterialId": _clean(payload.get("sourceMaterialId"), 40),
        "sourceDraftId": _clean(payload.get("sourceDraftId"), 40),
        "sourceJobId": _clean(payload.get("sourceJobId"), 40),
        "revisionNumber": payload.get("revisionNumber") or 1,
        "provenanceSha256": hashlib.sha256(full.encode("utf-8")).hexdigest()[:16],
    }
    # With bounded identifiers this remains below the OOXML/core-properties
    # 255-character limit enforced by python-pptx 1.x.
    return json.dumps(compact, ensure_ascii=False, separators=(",", ":"))


def _slide_header(line: str):
    patterns = (
        r"^\s*(?:#{1,3}\s*)?第\s*(\d{1,2})\s*(?:張|頁)\s*[：:.\-、]?\s*(.*)$",
        r"^\s*(?:#{1,3}\s*)?(\d{1,2})\s*(?:張|頁)\s*[：:.\-、]?\s*(.*)$",
        r"^\s*(?:#{1,3}\s*)?slide\s*(\d{1,2})\s*[：:.\-]?\s*(.*)$",
    )
    for pattern in patterns:
        match = re.match(pattern, line, re.I)
        if match:
            return int(match.group(1)), match.group(2).strip()
    return None


def parse_slide_outline(body: str, *, fallback_title: str = "AI 教學投影片") -> list[dict]:
    slides, current = [], None

    def push():
        nonlocal current
        if not current:
            return
        title = _clean(current.get("title") or f"第 {len(slides)+1} 張", 180)
        bullets = [_clean(item, 800) for item in current.get("bullets", []) if _clean(item, 800)][:_MAX_BULLETS]
        slides.append({"id": f"s{len(slides)+1}", "order": len(slides)+1, "enabled": True,
                       "title": title or f"第 {len(slides)+1} 張", "bullets": bullets, "speakerNotes": ""})
        current = None

    for raw in str(body or "").splitlines():
        raw = raw.strip()
        if not raw:
            continue
        header = _slide_header(raw)
        if header:
            push(); current = {"title": header[1] or f"第 {header[0]} 張", "bullets": []}; continue
        item = re.sub(r"^\s*(?:[-*•▪◦]|\d+[.)、])\s*", "", raw).strip()
        if not item:
            continue
        if current is None:
            current = {"title": fallback_title if not slides else f"第 {len(slides)+1} 張", "bullets": []}
        current["bullets"].append(item)
        if len(current["bullets"]) >= _MAX_BULLETS:
            push()
    push()
    if not slides:
        raise ValueError("AI 投影片草稿沒有可解析的內容。")
    if len(slides) > _MAX_SLIDES:
        raise ValueError(f"AI 投影片草稿超過 {_MAX_SLIDES} 張；請先分拆教材，系統不會靜默截斷。")
    return slides


def normalize_slides(slides: list[dict]) -> list[dict]:
    raw_slides = list(slides or [])
    if len(raw_slides) > _MAX_SLIDES:
        raise ValueError(f"PowerPoint 最多接受 {_MAX_SLIDES} 張來源投影片；請先分拆教材，系統不會靜默截斷。")
    normalized = []
    for index, raw in enumerate(raw_slides, 1):
        if not isinstance(raw, dict):
            continue
        layout = str(raw.get("layout") or raw.get("layoutKind") or "content").strip().lower()
        if layout not in _LAYOUT_KINDS:
            layout = "content"
        blocks = _normalize_blocks(raw.get("blocks"))
        raw_bullets = list(raw.get("bullets") or [])
        if len(raw_bullets) > _MAX_BULLETS:
            raise ValueError(f"單張投影片最多接受 {_MAX_BULLETS} 個 bullet；請分拆內容，系統不會靜默截斷。")
        normalized.append({
            "id": _clean(raw.get("id") or f"s{index}", 80), "order": index,
            "enabled": bool(raw.get("enabled", True)), "title": _clean(raw.get("title"), 180) or f"第 {index} 張",
            "bullets": [_clean(v, 800) for v in raw_bullets if _clean(v, 800)],
            "layout": layout, "blocks": blocks, "speakerNotes": _clean(raw.get("speakerNotes"), 4000),
        })
    if not normalized or not any(item["enabled"] for item in normalized):
        raise ValueError("至少要保留一張啟用的投影片。")
    return normalized


def _normalize_blocks(raw_blocks: Any) -> list[dict]:
    """Normalize declarative, safe blocks. URLs and HTML are deliberately absent."""
    raw_list = list(raw_blocks or [])
    if len(raw_list) > 12:
        raise ValueError("單張投影片最多接受 12 個內容區塊；請分拆投影片，系統不會靜默截斷。")
    blocks = []
    for raw in raw_list:
        if not isinstance(raw, dict):
            continue
        kind = str(raw.get("type") or "").strip().lower()
        if kind not in _BLOCK_KINDS:
            continue
        block = {"type": kind, "title": _clean(raw.get("title"), 140), "caption": _clean(raw.get("caption"), 300)}
        if kind == "image":
            asset = raw.get("asset") if isinstance(raw.get("asset"), dict) else {}
            backend = _clean(asset.get("backend"), 20).lower()
            key = _clean(asset.get("key"), 260)
            digest = _clean(asset.get("sha256"), 64).lower()
            mime = _clean(asset.get("mimeType"), 80).lower()
            if backend in {"r2", "oci", "gdrive", "mega"} and _IMAGE_KEY.fullmatch(key) and re.fullmatch(r"[0-9a-f]{64}", digest or "") and mime in _IMAGE_MIME_TYPES:
                block["asset"] = {"backend": backend, "key": key, "sha256": digest, "mimeType": mime}
            block["altText"] = _clean(raw.get("altText"), 300)
            block["fit"] = str(raw.get("fit") or "contain").strip().lower() if str(raw.get("fit") or "contain").strip().lower() in {"contain", "crop"} else "contain"
            source_label = _clean(raw.get("sourceLabel"), 180)
            lowered_source = source_label.lower()
            if (_IMAGE_KEY.fullmatch(source_label) or _LOCAL_PATH_PATTERN.search(source_label)
                    or lowered_source.startswith(("http://", "https://", "file://"))
                    or any(marker in lowered_source for marker in _SECRET_MARKERS)):
                source_label = ""
            block["sourceLabel"] = source_label
        elif kind == "chart":
            chart_type = _clean(raw.get("chartType") or "bar", 20).lower()
            labels = [_clean(v, 60) for v in list(raw.get("labels") or [])[:12] if _clean(v, 60)]
            values = []
            for value in list(raw.get("values") or [])[:12]:
                try: values.append(max(-1000000, min(1000000, float(value))))
                except (TypeError, ValueError): values.append(0)
            block.update({"chartType": chart_type if chart_type in {"bar", "column", "line"} else "bar", "labels": labels, "values": values[:len(labels)]})
        elif kind == "table":
            raw_headers = list(raw.get("headers") or [])
            raw_rows = list(raw.get("rows") or [])
            if len(raw_headers) > 8 or len(raw_rows) > 200:
                raise ValueError("表格最多接受 8 欄、200 列；請分拆資料，系統不會靜默截斷。")
            headers = [_clean(v, 80) for v in raw_headers if _clean(v, 80)]
            rows = []
            for row in raw_rows:
                if isinstance(row, list):
                    if len(row) > len(headers):
                        raise ValueError("表格資料欄數超過表頭欄數，拒絕靜默截斷。")
                    rows.append([_clean(v, 220) for v in row])
            block.update({"headers": headers, "rows": rows})
        elif kind == "comparison":
            left_items = list(raw.get("leftItems") or []); right_items = list(raw.get("rightItems") or [])
            if len(left_items) > 100 or len(right_items) > 100:
                raise ValueError("比較內容單側最多接受 100 項；請分拆資料，系統不會靜默截斷。")
            block.update({"leftTitle": _clean(raw.get("leftTitle"), 100), "rightTitle": _clean(raw.get("rightTitle"), 100),
                          "leftItems": [_clean(v, 220) for v in left_items if _clean(v, 220)],
                          "rightItems": [_clean(v, 220) for v in right_items if _clean(v, 220)]})
        else:
            block["text"] = _clean(raw.get("text"), 1200)
        blocks.append(block)
    return blocks


def _write_notes(slide, text: str) -> None:
    try:
        slide.notes_slide.notes_text_frame.text = text
    except Exception:
        pass


def _clear_template_slides(prs) -> None:
    """Keep the template theme/layouts but remove authoring/sample slide instances."""
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.slides._sldIdLst.remove(slide_id)
        try:
            prs.part.drop_rel(rel_id)
        except KeyError:
            pass


def _layout_for(prs, kind: str, profile: dict | None = None, *, quality_report=None, slide_id=""):
    """Resolve a template layout by safe profile, friendly name, then deterministic fallback."""
    layouts = prs.slide_layouts
    if not layouts:
        raise RuntimeError("PowerPoint 範本沒有可用 layout。")
    profile_map = dict((profile or {}).get("layoutMap") or {})
    desired = str(profile_map.get(kind) or "").strip().lower()
    aliases = {
        "title": ("title slide", "title"), "section": ("section header", "section", "title and content"),
        "content": ("title and content", "content"), "image": ("picture", "title only", "blank"),
        "comparison": ("comparison", "two content", "title and content"), "table": ("title and content", "comparison"),
        "summary": ("title and content", "section header"),
    }
    candidates = ((desired,) if desired else ()) + aliases.get(kind, aliases["content"])
    for candidate in candidates:
        for layout in layouts:
            if candidate and candidate in str(layout.name or "").lower():
                return layout
    preferred = 0 if kind == "title" else (1 if len(layouts) > 1 else 0)
    if isinstance(quality_report, dict):
        updated = quality.add_warning(quality_report, "TEMPLATE_FALLBACK", slide_id=slide_id,
                                      detail=f"範本缺少 {kind} layout，已使用安全 fallback。")
        quality_report.clear(); quality_report.update(updated)
    return layouts[preferred]


def _placeholder_score(shape, role: str) -> int:
    name = str(getattr(shape, "name", "") or "").lower()
    score = 0
    if role == "body" and any(token in name for token in ("content", "body", "text", "內容", "文字")):
        score += 4
    if role == "image" and any(token in name for token in ("picture", "image", "photo", "圖片", "影像")):
        score += 4
    if role == "table" and any(token in name for token in ("table", "grid", "表格")):
        score += 4
    if hasattr(shape, "text_frame"):
        score += 1
    return score


def _body_placeholder(slide, profile: dict | None = None):
    candidates = [p for p in slide.placeholders if p != slide.shapes.title and hasattr(p, "text_frame")]
    if not candidates:
        return None
    selector = str(((profile or {}).get("placeholderMap") or {}).get("body") or "auto").strip().lower()
    if selector != "auto":
        named = [shape for shape in candidates if selector in str(getattr(shape, "name", "") or "").lower()]
        if named:
            return named[0]
    candidates.sort(key=lambda shape: _placeholder_score(shape, "body"), reverse=True)
    return candidates[0]


def _placeholder_box(slide, role: str, profile: dict | None = None):
    """Return a safe geometry box from a matching template placeholder, if one exists."""
    candidates = [p for p in slide.placeholders if p != slide.shapes.title]
    selector = str(((profile or {}).get("placeholderMap") or {}).get(role) or "auto").strip().lower()
    if selector != "auto":
        selected = [shape for shape in candidates if selector in str(getattr(shape, "name", "") or "").lower()]
        if selected:
            candidates = selected
    candidates.sort(key=lambda shape: _placeholder_score(shape, role), reverse=True)
    if not candidates or _placeholder_score(candidates[0], role) <= 0:
        return None
    shape = candidates[0]
    emu = float(Inches(1))
    try:
        left, top, width, height = [float(value) / emu for value in (shape.left, shape.top, shape.width, shape.height)]
    except Exception:
        return None
    if width < 1.0 or height < 0.6:
        return None
    return (left, top, width, height)


def _textbox(slide, left, top, width, height, text=""):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
    return shape.text_frame


def _placeholder_block(slide, block: dict, message: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.55))
    shape.text_frame.text = "素材待補\n" + _clean(block.get("title") or block.get("altText") or message, 220)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(14)


def _set_body_font(frame, bullets: list[str], max_size: int = 24) -> None:
    total_chars = sum(len(str(value or "")) for value in bullets)
    size = 24 if total_chars <= 360 else (22 if total_chars <= 560 else 20)
    size = max(18, min(size, max_size))
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(size)


def _render_table(slide, block: dict, profile: dict | None = None) -> None:
    headers, rows = block.get("headers") or [], block.get("rows") or []
    if not headers:
        _placeholder_block(slide, block, "表格欄位不足"); return
    box = _placeholder_box(slide, "table", profile) or (0.75, 2.0, 11.8, 4.7)
    table = slide.shapes.add_table(len(rows) + 1, len(headers), *(Inches(value) for value in box)).table
    for col, text in enumerate(headers):
        table.cell(0, col).text = text
        for p in table.cell(0, col).text_frame.paragraphs: p.font.size = Pt(16)
    for row_idx, row in enumerate(rows, 1):
        for col, text in enumerate(row):
            table.cell(row_idx, col).text = text
            for p in table.cell(row_idx, col).text_frame.paragraphs: p.font.size = Pt(14)


def _render_chart(slide, block: dict) -> None:
    labels, values = block.get("labels") or [], block.get("values") or []
    if not labels or len(values) != len(labels):
        _placeholder_block(slide, block, "圖表資料不足"); return
    data = CategoryChartData(); data.categories = labels; data.add_series(block.get("title") or "數值", values)
    chart_type = {"line": XL_CHART_TYPE.LINE_MARKERS, "column": XL_CHART_TYPE.COLUMN_CLUSTERED}.get(block.get("chartType"), XL_CHART_TYPE.BAR_CLUSTERED)
    chart = slide.shapes.add_chart(chart_type, Inches(0.75), Inches(2.0), Inches(11.8), Inches(4.7), data).chart
    chart.has_legend = False


def _render_comparison(slide, block: dict) -> None:
    for left, title, items in ((0.75, block.get("leftTitle"), block.get("leftItems") or []), (6.75, block.get("rightTitle"), block.get("rightItems") or [])):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(2.0), Inches(5.1), Inches(4.5))
        shape.text_frame.text = _clean(title or "比較項目", 100) + "\n" + "\n".join("• " + value for value in items)
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(18)


def _image_box(path: Path, *, fit: str, left=1.0, top=1.7, width=10.9, height=4.55):
    """Return safe image placement preserving aspect ratio; crop uses a local temporary derivative."""
    if Image is None:
        return path, left, top, width, None
    with Image.open(path) as image:
        iw, ih = image.size
        if not iw or not ih:
            return path, left, top, width, None
        image_ratio = iw / ih
        box_ratio = width / height
        if fit == "crop":
            if image_ratio > box_ratio:
                target_w = int(ih * box_ratio); x0 = max(0, (iw - target_w) // 2); crop = (x0, 0, x0 + target_w, ih)
            else:
                target_h = int(iw / box_ratio); y0 = max(0, (ih - target_h) // 2); crop = (0, y0, iw, y0 + target_h)
            derived = path.with_name(path.stem + "-crop" + path.suffix)
            image.crop(crop).save(derived)
            return derived, left, top, width, height
        if image_ratio >= box_ratio:
            rendered_w = width; rendered_h = width / image_ratio
            return path, left, top + (height - rendered_h) / 2, rendered_w, rendered_h
        rendered_h = height; rendered_w = height * image_ratio
        return path, left + (width - rendered_w) / 2, top, rendered_w, rendered_h


def _side_image_box(slide_width: float) -> tuple[float, float, float, float]:
    """Right-hand picture area used when a slide also has bullets (text stays on the left)."""
    width = max(3.0, slide_width * 0.43)
    return (slide_width * 0.53, 1.7, width, 4.55)


def _full_image_box(slide_width: float) -> tuple[float, float, float, float]:
    margin = max(0.5, slide_width * 0.075)
    return (margin, 1.7, max(3.0, slide_width - 2 * margin), 4.55)


def _picture_placeholder_box(slide, profile: dict | None = None):
    """Geometry of a template placeholder that is really meant for pictures (else None).

    A body/text placeholder is never a picture area, otherwise a picture would be
    drawn over the slide's own bullets.
    """
    selector = str(((profile or {}).get("placeholderMap") or {}).get("image") or "auto").strip().lower()
    named = selector != "auto" and any(
        selector in str(getattr(shape, "name", "") or "").lower() for shape in slide.placeholders if shape != slide.shapes.title)
    if not named and not any(
            _placeholder_score(shape, "image") >= 4 for shape in slide.placeholders if shape != slide.shapes.title):
        return None
    return _placeholder_box(slide, "image", profile)


def _render_image(slide, block: dict, image_resolver=None, profile: dict | None = None,
                  image_box: tuple[float, float, float, float] | None = None) -> None:
    asset = block.get("asset") if isinstance(block.get("asset"), dict) else None
    path = None
    try:
        path = image_resolver(asset) if asset and callable(image_resolver) else None
        if path and Path(path).is_file():
            box = image_box or (1.0, 1.7, 10.9, 4.55)
            image_path, left, top, width, height = _image_box(Path(path), fit=str(block.get("fit") or "contain"),
                                                               left=box[0], top=box[1], width=box[2], height=box[3])
            kwargs = {"width": Inches(width)}
            if height is not None: kwargs["height"] = Inches(height)
            slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), **kwargs)
            caption = _clean(block.get("caption"), 300)
            source_label = _clean(block.get("sourceLabel"), 180)
            visible = "｜".join(value for value in (caption, source_label) if value)
            if visible:
                frame = _textbox(slide, box[0], 6.35, box[2], 0.45, visible)
                for paragraph in frame.paragraphs: paragraph.font.size = Pt(10)
            return
    except Exception:
        # A missing/invalid allowed asset is presentation content, not a worker failure.
        pass
    _placeholder_block(slide, block, "允許的圖片素材尚未提供")


def _render_blocks(slide, blocks: list[dict], image_resolver=None, profile: dict | None = None,
                   image_box: tuple[float, float, float, float] | None = None) -> None:
    for block in blocks:
        kind = block.get("type")
        if kind == "image": _render_image(slide, block, image_resolver, profile, image_box)
        elif kind == "chart": _render_chart(slide, block)
        elif kind == "table": _render_table(slide, block, profile)
        elif kind == "comparison": _render_comparison(slide, block)
        elif kind == "callout":
            frame = _textbox(slide, 0.9, 5.8, 11.4, 0.7, block.get("text") or block.get("caption") or "")
            for paragraph in frame.paragraphs: paragraph.font.size = Pt(16)


def _branding_footer(slide, branding: dict[str, Any] | None) -> None:
    if not branding:
        return
    parts = []
    for key in ("groupLabel", "areaLabel", "teacherName"):
        value = _clean((branding or {}).get(key), 80)
        if value: parts.append(value)
    revision = str((branding or {}).get("revisionNumber") or "").strip()
    rendered_date = _clean((branding or {}).get("renderedDate") or (branding or {}).get("publishedDate"), 40)
    if revision: parts.append(f"r{revision}")
    if rendered_date: parts.append(rendered_date)
    if not parts:
        return
    frame = _textbox(slide, 0.6, 7.05, 12.0, 0.25, " · ".join(parts))
    for paragraph in frame.paragraphs: paragraph.font.size = Pt(8)


def _with_formal_cover(title: str, slides: list[dict]) -> list[dict]:
    """Prepend a deterministic title cover when authored content does not already start with one."""
    prepared = list(slides or [])
    if prepared and str(prepared[0].get("layout") or "").strip().lower() == "title":
        return prepared
    cover = {
        "id": "phase4-cover",
        "order": 1,
        "enabled": True,
        "title": _clean(title, 180) or "AI 教學投影片",
        "bullets": [],
        "layout": "title",
        "blocks": [],
        "speakerNotes": "",
    }
    output = [cover, *prepared]
    for index, item in enumerate(output, 1):
        item["order"] = index
    return output


def render_pptx(*, title: str, slides: list[dict], output_path: Path,
                provenance: dict[str, Any], template_path: Path | None = None,
                layout_profile: dict | None = None, image_resolver=None,
                branding: dict[str, Any] | None = None, quality_report: dict[str, Any] | None = None) -> Path:
    if Presentation is None or Inches is None:
        raise RuntimeError("AI Worker 尚未安裝 python-pptx；請更新 requirements 後重新啟動。")
    safe_provenance = repository.sanitize_provenance(provenance)
    prepared, manifest = quality.prepare_slides(
        normalize_slides(slides),
        provenance_present=bool(safe_provenance.get("sourceMaterialId") and safe_provenance.get("sourceDraftId")),
    )
    prepared = _with_formal_cover(title, prepared)
    prs = Presentation(str(template_path)) if template_path else Presentation()
    if template_path:
        _clear_template_slides(prs)
    slide_width_in = float(prs.slide_width or Inches(10)) / float(Inches(1))
    core = prs.core_properties
    core.title = _clean(title, 255)
    core.subject = "Teacher AI reviewed teaching presentation"
    core.keywords = f"Teacher,AI,medical-laboratory,review-required,{quality.RULESET_VERSION}"
    provenance_json = json.dumps(safe_provenance, ensure_ascii=False, separators=(",", ":"))
    core.comments = _core_provenance_comment(safe_provenance)
    for item in prepared:
        slide = prs.slides.add_slide(_layout_for(prs, item.get("layout") or "content", layout_profile,
                                                 quality_report=manifest, slide_id=item.get("id") or ""))
        if slide.shapes.title is not None:
            slide.shapes.title.text = item["title"]
        body = _body_placeholder(slide, layout_profile)
        if body is None:
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8))
        # A slide with bullets AND a picture keeps its text on the left and the
        # picture on the right instead of drawing the picture over the text.
        has_image = any(isinstance(b, dict) and b.get("type") == "image" for b in (item.get("blocks") or []))
        image_box = _picture_placeholder_box(slide, layout_profile) if has_image else None
        side_box = None
        if has_image and image_box is None:
            if item["bullets"]:
                image_box = side_box = _side_image_box(slide_width_in)
            else:
                image_box = _full_image_box(slide_width_in)
        if side_box is not None:
            try:
                left, top, width, height = body.left, body.top, body.width, body.height
                limit = Inches(max(3.0, side_box[0] - 0.3)) - left
                if width > limit > 0:
                    body.left, body.top, body.width, body.height = left, top, limit, height
            except Exception:
                pass  # geometry is best effort; the picture still renders in its own box
        frame = body.text_frame; frame.clear()
        for index, bullet in enumerate(item["bullets"]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet; paragraph.level = 0
        _set_body_font(frame, item["bullets"], max_size=20 if side_box else 24)
        _render_blocks(slide, item.get("blocks") or [], image_resolver=image_resolver, profile=layout_profile, image_box=image_box)
        _branding_footer(slide, branding)
        notes = "\n\n".join(filter(None, [_clean(item.get("speakerNotes"), 4000), "PROVENANCE " + provenance_json]))
        _write_notes(slide, notes)
    manifest["slideCount"] = len(prepared)
    manifest = quality.sanitize_quality_manifest(manifest)
    if isinstance(quality_report, dict):
        quality_report.clear(); quality_report.update(manifest)
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True); prs.save(str(output_path))
    if not output_path.is_file() or output_path.stat().st_size < 1024:
        raise RuntimeError("python-pptx 未產生有效 PowerPoint 檔案。")
    return output_path


def _template_location(item: dict) -> dict:
    return {"backend": item.get("storageBackend"), "key": item.get("storageKey"), "sha256": item.get("sha256")}


def _source_context(draft: dict, source: dict, *, template_id: str = "", revision_number: int = 1) -> dict[str, Any]:
    chunk_ids = []
    for chunk in list(draft.get("sourceChunks") or []):
        value = str((chunk or {}).get("chunkId") or (chunk or {}).get("id") or "") if isinstance(chunk, dict) else str(chunk or "")
        if value.strip():
            chunk_ids.append(value.strip())
    return {
        "sourceMaterialId": str(source.get("id") or ""),
        "sourceMaterialVersion": max(1, int(source.get("currentVersion") or 1)),
        "sourceDraftId": str(draft.get("id") or ""),
        "sourceJobId": str(draft.get("sourceJobId") or ""), "sourceChunkIds": chunk_ids,
        "provider": str(draft.get("provider") or ""), "model": str(draft.get("model") or ""),
        "templateId": str(template_id or ""), "teacherApprovedBy": str(draft.get("approvedBy") or ""),
        "teacherApprovedAt": str(draft.get("approvedAt") or ""), "revisionNumber": max(1, int(revision_number or 1)),
    }


def _branding_context(*, title: str, group: str, area: str, teacher: str, revision: int) -> dict[str, Any]:
    return {
        "title": _clean(title, 180),
        "groupLabel": _clean(group, 80),
        "areaLabel": _clean(area, 80),
        "teacherName": _clean(teacher, 80),
        "revisionNumber": max(1, int(revision or 1)),
        "renderedDate": time.strftime("%Y-%m-%d", time.gmtime()),
    }


def _validated_template(template_id: str, *, group: str, area: str):
    template = repository.get_template(template_id) if template_id else None
    if template_id and (not template or not template.get("active") or template.get("group") != group or template.get("area") != area):
        raise RuntimeError("PowerPoint 範本不存在、已停用或超出允許範圍。")
    return template


def _image_resolver(storage: PresentationStorage, root: Path, local: dict[str, Path] | None = None):
    """Resolve only normalized shared-provider image assets; failures become slide fallbacks.

    ``local`` maps sha256 -> a picture this job just extracted, so those need no
    second download from the shared provider.
    """
    counter = 0
    def resolve(asset):
        nonlocal counter
        if not isinstance(asset, dict):
            return None
        key = str(asset.get("key") or "")
        if not _IMAGE_KEY.fullmatch(key):
            return None
        cached = (local or {}).get(str(asset.get("sha256") or "").lower())
        if cached and Path(cached).is_file():
            return Path(cached)
        counter += 1
        suffix = ".png" if str(asset.get("mimeType") or "") == "image/png" else ".jpg"
        try:
            return storage.download(asset, root / f"image-{counter}{suffix}")
        except Exception:
            # A stale legacy provider reference must not make the whole
            # PowerPoint job fail. The slide renderer already has a visual
            # fallback for a missing optional image.
            return None
    return resolve


def _auto_pictures_enabled(job: dict) -> bool:
    if str(os.environ.get("AI_PRESENTATION_AUTO_PICTURES", "true")).strip().lower() in {"0", "false", "no", "off"}:
        return False
    return dict(job.get("request") or {}).get("autoPictures") is not False


def _picture_sources(draft: dict, source: dict) -> list[dict]:
    """The material the deck was written from, plus the reference files the teacher added."""
    materials = [source]
    job_id = str(draft.get("sourceJobId") or "")
    try:
        script_job = media_script_repository.get_job(job_id) if job_id else None
        for reference_id in list(((script_job or {}).get("request") or {}).get("referenceMaterialIds") or [])[:9]:
            reference = material_repository.get_material(str(reference_id or ""))
            if (reference and str(reference.get("id") or "") != str(source.get("id") or "")
                    and reference.get("group") == source.get("group") and reference.get("area") == source.get("area")):
                materials.append(reference)
    except Exception:
        pass  # references are optional; the main material is still scanned
    return materials


def _attach_source_pictures(slides: list[dict], *, draft: dict, source: dict, storage: PresentationStorage,
                            workdir: Path, local: dict[str, Path]) -> tuple[list[dict], dict[str, Any]]:
    """Give slides pictures found in the teacher's own sources. Never fails the job."""
    try:
        if not storage.image_backend():
            return slides, {"unsupported": True}
        from teacher_app.assessments import ai_runtime

        def store(path: Path, sha: str, mime: str) -> dict[str, Any]:
            asset = storage.store_image(path, sha256=sha, mime_type=mime)
            local[sha] = path
            return asset

        return picture_picker.attach_pictures(
            slides, _picture_sources(draft, source), workdir=workdir,
            fetch_source=lambda material: ai_runtime.material_source_to_temp(material), store_image=store)
    except Exception:
        LOGGER.warning("AI slide auto pictures skipped", exc_info=True)
        return slides, {"failed": True}


def generate_presentation(*, job: dict, progress_callback=None, storage: PresentationStorage | None = None) -> dict:
    storage = storage or PresentationStorage()
    existing = repository.get_presentation_by_source_job_id(str(job.get("id") or ""))
    if existing:
        return {"presentationId": existing["id"], "replayed": True}
    draft = media_script_repository.get_script(str(job.get("draftId") or ""))
    if not draft or draft.get("draftType") != "slides" or draft.get("status") != "approved":
        raise RuntimeError("只有已由授課教師核准的投影片大綱可產生 PowerPoint。")
    if not str(draft.get("approvedBy") or "") or not str(draft.get("approvedAt") or ""):
        raise RuntimeError("AI 投影片大綱缺少教師核准紀錄。")
    if str(draft.get("group") or "") != str(job.get("group") or "") or str(draft.get("area") or "") != str(job.get("area") or ""):
        raise RuntimeError("PowerPoint 工作範圍與來源草稿不一致。")
    source = material_repository.get_material(str(draft.get("materialId") or ""))
    if not source or source.get("group") != draft.get("group") or source.get("area") != draft.get("area"):
        raise RuntimeError("來源教材不存在或範圍與 AI 草稿不一致。")
    template_id = str(job.get("templateId") or "")
    template = _validated_template(template_id, group=str(draft.get("group") or ""), area=str(draft.get("area") or ""))
    if progress_callback:
        progress_callback(15, "解析投影片大綱", "套用 Phase 4 智慧版型與品質規則")
    supplied = dict(job.get("request") or {}).get("slides")
    slides = normalize_slides(supplied) if isinstance(supplied, list) else parse_slide_outline(str(draft.get("body") or ""), fallback_title=str(draft.get("title") or "AI 教學投影片"))
    # 老師不必自己找切點：同一份教材若已有「已核准講稿」，且投影片還沒有任何講者備註，
    # 就依內容把講稿自動分段，寫進每頁備註（AI 影片會逐頁念備註）。失敗不影響產檔。
    try:
        approved_script = script_alignment.latest_approved_script(media_script_repository.list_scripts(str(source.get("id") or "")))
        if approved_script:
            slides, notes_info = script_alignment.attach_script_notes(slides, str(approved_script.get("body") or ""))
            if notes_info.get("applied") and progress_callback:
                progress_callback(20, "依講稿自動分段", f"已把已核准講稿的 {notes_info['unitCount']} 個段落對應到 {notes_info['slideCount']} 張投影片備註")
    except Exception:
        pass
    started = time.perf_counter(); quality_report: dict[str, Any] = {}
    with tempfile.TemporaryDirectory(prefix="teacher-ppt-") as temp:
        root = Path(temp); template_path = None
        template_fallback_reason = ""
        # 老師不必自己找圖：從教材與參考資料（Word／PowerPoint／PDF／圖片檔）抽出圖片，
        # 依周圍文字自動配到最相關的投影片。失敗或沒有合適的圖都不影響產檔。
        picture_info: dict[str, Any] = {}
        local_pictures: dict[str, Path] = {}
        if _auto_pictures_enabled(job):
            if progress_callback: progress_callback(25, "從教材挑選圖片", "掃描教材與參考資料裡的圖片，配到相關的投影片")
            slides, picture_info = _attach_source_pictures(slides, draft=draft, source=source, storage=storage,
                                                           workdir=root / "pictures", local=local_pictures)
            if picture_info.get("applied") and progress_callback:
                progress_callback(28, "已自動配圖", f"在 {picture_info['applied']} 張投影片放入教材中的相關圖片")
        if template:
            if progress_callback: progress_callback(30, "下載簡報範本", "從共享 provider 下載組別範本")
            try:
                template_path = storage.download(_template_location(template), root / "template.pptx")
            except Exception as exc:
                # Old group templates may still point at MEGA even after the
                # production artifact store moved to R2. Keep generation
                # available by using the built-in safe layout rather than
                # failing the entire job on a legacy template login.
                template_path = None
                template_fallback_reason = str(exc)[:240]
                if progress_callback:
                    progress_callback(35, "範本改用安全預設版型", "舊範本目前無法讀取；PowerPoint 仍會繼續產生。")
        output = root / "presentation.pptx"
        if progress_callback: progress_callback(55, "建立 PowerPoint", "自動拆頁、圖片適配並寫入 speaker notes / provenance")
        render_pptx(title=str(draft.get("title") or source.get("title") or "AI 教學投影片"), slides=slides,
                    output_path=output, provenance=_source_context(draft, source, template_id=template_id), template_path=template_path,
                    layout_profile=(template or {}).get("layoutProfile"), image_resolver=_image_resolver(storage, root, local_pictures),
                    branding=_branding_context(title=str(draft.get("title") or source.get("title") or "AI 教學投影片"),
                                               group=str(draft.get("group") or ""), area=str(draft.get("area") or ""),
                                               teacher=str(draft.get("approvedBy") or ""), revision=1),
                    quality_report=quality_report)
        if picture_info.get("applied"):
            # 圖片是自動挑的：請授課教師確認內容相符、沒有病人個資後再發布。
            detail =f"已自動從教材配上 {picture_info['applied']} 張圖片，請確認圖片內容相符且不含病人個資。"
            updated = quality.add_warning(quality_report, "AUTO_PICTURES", detail=detail)
            quality_report.clear(); quality_report.update(updated)
        if template_fallback_reason:
            warnings=list(quality_report.get("warnings") or [])
            warnings.append({"code":"TEMPLATE_PROVIDER_FALLBACK","message":"組別範本無法從舊儲存讀取，已改用安全預設版型。"})
            quality_report["warnings"]=warnings
        if progress_callback: progress_callback(80, "保存 PowerPoint", "將產出檔保存至共享 provider")
        artifact = storage.store(output, namespace="artifacts", object_id=str(job.get("id") or ""), filename=f"{str(draft.get('title') or 'AI教學投影片')[:60]}.pptx")
    provenance = repository.sanitize_provenance(_source_context(draft, source, template_id=template_id))
    metrics = quality.sanitize_render_metrics({
        "rulesetVersion": quality.RULESET_VERSION,
        "durationMs": round((time.perf_counter() - started) * 1000),
        "attempts": int(job.get("attempts") or 0),
        "jobId": str(job.get("id") or ""),
        "templateFallbackCount": sum(1 for item in quality_report.get("warnings", []) if item.get("code") == "TEMPLATE_FALLBACK"),
        "renderedSlideCount": int(quality_report.get("slideCount") or 0),
    })
    created = repository.create_presentation(
        material_id=str(source.get("id") or ""), draft_id=str(draft.get("id") or ""), template_id=template_id,
        group_key=str(draft.get("group") or ""), training_area=str(draft.get("area") or ""),
        title=str(draft.get("title") or source.get("title") or "AI 教學投影片")[:255], slides=slides,
        actor_username=str(job.get("actorUsername") or ""), source_job_id=str(job.get("id") or ""),
        provider=str(draft.get("provider") or ""), model=str(draft.get("model") or ""), artifact=artifact,
        provenance=provenance, quality_manifest=quality_report, render_metrics=metrics,
        render_ruleset_version=quality.RULESET_VERSION)
    if progress_callback: progress_callback(95, "PowerPoint 已保存", "品質檢查完成；等待授課教師檢查、編修與核准")
    return {"presentationId": str(created.get("id") or ""), "artifactSha256": str(created.get("artifactSha256") or ""),
            "artifactBytes": int(created.get("artifactBytes") or 0), "qualityStatus": str((created.get("qualityManifest") or {}).get("status") or ""),
            "replayed": False}


def generate_revision(*, job: dict, progress_callback=None, storage: PresentationStorage | None = None) -> dict:
    """Worker-only render path for an already-created immutable draft revision."""
    storage = storage or PresentationStorage(); request_payload = dict(job.get("request") or {})
    presentation_id = str(request_payload.get("presentationId") or "").strip()
    current = repository.get_presentation(presentation_id)
    if not current:
        raise RuntimeError("找不到待重新產生的 PowerPoint revision。")
    if current.get("status") != "draft":
        raise RuntimeError("只有 draft PowerPoint revision 可以重新產生 artifact。")
    if current.get("artifactStorageKey") and current.get("artifactSha256") and int(current.get("artifactBytes") or 0) > 0:
        return {"presentationId": presentation_id, "artifactSha256": current.get("artifactSha256"), "artifactBytes": current.get("artifactBytes"), "replayed": True}
    if current.get("group") != job.get("group") or current.get("area") != job.get("area"):
        raise RuntimeError("PowerPoint revision 工作範圍不一致。")
    draft = media_script_repository.get_script(str(current.get("draftId") or "")) or {}
    source = material_repository.get_material(str(current.get("materialId") or "")) or {}
    if (
        not draft
        or not source
        or draft.get("group") != current.get("group")
        or source.get("group") != current.get("group")
        or draft.get("area") != current.get("area")
        or source.get("area") != current.get("area")
    ):
        raise RuntimeError("PowerPoint revision 來源不存在或授權範圍已改變。")
    template_id = str(current.get("templateId") or "")
    template = _validated_template(template_id, group=str(current.get("group") or ""), area=str(current.get("area") or ""))
    slides = normalize_slides(list(current.get("slides") or [])); started = time.perf_counter(); quality_report: dict[str, Any] = {}
    if progress_callback: progress_callback(25, "準備 PowerPoint revision", "套用 Phase 4 智慧版型與品質規則")
    with tempfile.TemporaryDirectory(prefix="teacher-ppt-revision-") as temp:
        root = Path(temp); template_path = None
        template_fallback_reason = ""
        if template:
            try:
                template_path = storage.download(_template_location(template), root / "template.pptx")
            except Exception as exc:
                template_path = None
                template_fallback_reason = str(exc)[:240]
                if progress_callback:
                    progress_callback(35, "範本改用安全預設版型", "舊範本目前無法讀取；revision 仍會繼續產生。")
        output = root / "presentation.pptx"
        if progress_callback: progress_callback(55, "建立 PowerPoint revision", "自動拆頁、圖片適配並寫入安全 provenance")
        render_pptx(title=str(current.get("title") or "AI 教學投影片"), slides=slides, output_path=output,
                    provenance=_source_context(draft, source, template_id=template_id, revision_number=int(current.get("revisionNumber") or 1)), template_path=template_path,
                    layout_profile=(template or {}).get("layoutProfile"), image_resolver=_image_resolver(storage, root),
                    branding=_branding_context(title=str(current.get("title") or "AI 教學投影片"), group=str(current.get("group") or ""),
                                               area=str(current.get("area") or ""), teacher=str(draft.get("approvedBy") or current.get("updatedBy") or ""),
                                               revision=int(current.get("revisionNumber") or 1)),
                    quality_report=quality_report)
        if template_fallback_reason:
            warnings=list(quality_report.get("warnings") or [])
            warnings.append({"code":"TEMPLATE_PROVIDER_FALLBACK","message":"組別範本無法從舊儲存讀取，已改用安全預設版型。"})
            quality_report["warnings"]=warnings
        if progress_callback: progress_callback(80, "保存 PowerPoint revision", "將新 revision artifact 保存至共享 provider")
        artifact = storage.store(output, namespace="artifacts", object_id=presentation_id,
                                 filename=f"{str(current.get('title') or 'AI教學投影片')[:60]}-r{int(current.get('revisionNumber') or 1)}.pptx")
    updated = repository.update_presentation_artifact(presentation_id, artifact=artifact, actor_username=str(job.get("actorUsername") or ""))
    if not updated:
        raise RuntimeError("PowerPoint revision 已不存在。")
    metrics = quality.sanitize_render_metrics({
        "rulesetVersion": quality.RULESET_VERSION,
        "durationMs": round((time.perf_counter() - started) * 1000),
        "attempts": int(job.get("attempts") or 0),
        "jobId": str(job.get("id") or ""),
        "templateFallbackCount": sum(1 for item in quality_report.get("warnings", []) if item.get("code") == "TEMPLATE_FALLBACK"),
        "renderedSlideCount": int(quality_report.get("slideCount") or 0),
    })
    updated = repository.update_presentation_quality(
        presentation_id, quality_manifest=quality_report, render_metrics=metrics,
        render_ruleset_version=quality.RULESET_VERSION, actor_username=str(job.get("actorUsername") or ""),
    ) or updated
    if progress_callback: progress_callback(95, "PowerPoint revision 已保存", "品質檢查完成；等待授課教師檢查與核准")
    return {"presentationId": presentation_id, "artifactSha256": updated.get("artifactSha256"),
            "artifactBytes": int(updated.get("artifactBytes") or 0), "qualityStatus": str((updated.get("qualityManifest") or {}).get("status") or ""),
            "replayed": False}


__all__ = ["parse_slide_outline", "normalize_slides", "render_pptx", "generate_presentation", "generate_revision"]
